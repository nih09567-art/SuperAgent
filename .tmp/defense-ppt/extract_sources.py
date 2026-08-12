from __future__ import annotations

from pathlib import Path

from docx import Document
from pypdf import PdfReader


ROOT = Path(r"E:\Program\SuperAgent")
OUT = ROOT / ".tmp" / "defense-ppt"


def extract_docx(source: Path, output: Path) -> None:
    doc = Document(source)
    lines: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)
    for table_index, table in enumerate(doc.tables, start=1):
        lines.append(f"\n[TABLE {table_index}]")
        for row in table.rows:
            lines.append("\t".join(cell.text.strip().replace("\n", " / ") for cell in row.cells))
    output.write_text("\n".join(lines), encoding="utf-8")


def extract_pdf(source: Path, output: Path) -> None:
    reader = PdfReader(str(source))
    lines = [f"[PAGES] {len(reader.pages)}"]
    for index, page in enumerate(reader.pages, start=1):
        lines.append(f"\n===== PAGE {index} =====\n")
        lines.append(page.extract_text() or "")
    output.write_text("\n".join(lines), encoding="utf-8")


extract_pdf(
    Path(r"D:\xwechat_files\wxid_ddeayze1y1mm22_44d3\msg\file\2026-08\2025082916122218429(1).pdf"),
    OUT / "pdf" / "taskbook.txt",
)

extract_docx(
    Path(r"C:\Users\刘建杰\Desktop\多智能体协同与执行编排机制调研及原型成果报告.docx"),
    OUT / "docx" / "multi_agent_report.txt",
)

extract_docx(
    Path(r"C:\Users\刘建杰\Desktop\面向数字员工的主Agent意图识别路由与权限治理研究报告.docx"),
    OUT / "docx" / "main_agent_governance_report.txt",
)

