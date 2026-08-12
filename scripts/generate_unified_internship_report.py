"""Generate the integrated internship technical report from its Markdown source.

The output intentionally uses a single-column academic layout: A4 paper, numbered
sections, IEEE-style numeric references, Chinese body text in SimSun and Latin text
in Times New Roman.  The Markdown file remains the reviewable source of truth.
"""

from __future__ import annotations

import re
import shutil
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "面向数字员工的智能任务执行关键技术研究报告.md"
OUTPUT_DIR = ROOT / "output"
ASSET_DIR = OUTPUT_DIR / "unified_report_assets"
OUTPUT = OUTPUT_DIR / "面向数字员工的智能任务执行关键技术研究报告.docx"
EXTRACTED_MEDIA = ROOT / ".artifacts" / "report_a_media"
SCREENSHOT_DIR = Path(r"C:\Users\刘建杰\Pictures\Screenshots")

NAVY = "173A5E"
BLUE = "2B75A5"
TEAL = "2299A8"
GREEN = "347F63"
ORANGE = "CE7B29"
PURPLE = "725DAB"
RED = "B95151"
GRAY = "6E7D89"
LIGHT = "F3F6FA"
MID = "D7E2EC"
TEXT = "22313F"


def _available_font(preferred: list[str]) -> str:
    names = {font_manager.FontProperties(fname=p).get_name() for p in font_manager.findSystemFonts()}
    return next((name for name in preferred if name in names), "DejaVu Sans")


CN_FONT = _available_font(["Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "SimSun"])
plt.rcParams["font.sans-serif"] = [CN_FONT, "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False


def _hex(value: str) -> str:
    return f"#{value}"


def _box(ax, x: float, y: float, w: float, h: float, text: str, color: str, size: int = 11) -> None:
    patch = FancyBboxPatch(
        (x, y),
        w,
        h,
        boxstyle="round,pad=0.012,rounding_size=0.025",
        linewidth=1.2,
        edgecolor="white",
        facecolor=_hex(color),
    )
    ax.add_patch(patch)
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", color="white", fontsize=size, weight="bold")


def _arrow(ax, start: tuple[float, float], end: tuple[float, float], color: str = GRAY) -> None:
    ax.add_patch(FancyArrowPatch(start, end, arrowstyle="-|>", mutation_scale=14, linewidth=1.4, color=_hex(color)))


def _save(fig, name: str) -> None:
    fig.savefig(ASSET_DIR / name, dpi=240, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def make_overall_architecture() -> None:
    fig, ax = plt.subplots(figsize=(12, 6.5))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.text(0.5, 0.95, "SuperAgent 面向数字员工的四平面总体架构", ha="center", va="center", fontsize=18, weight="bold", color=_hex(NAVY))
    _box(ax, 0.04, 0.73, 0.92, 0.12, "决策平面：Intent Catalog → TaskProfile → AgentCard → RoutingDecision", BLUE, 12)
    _box(ax, 0.04, 0.54, 0.92, 0.12, "执行平面：候选计划校验 → TaskGraph → Scheduler → Agent / MCP Tool", PURPLE, 12)
    _box(ax, 0.04, 0.35, 0.92, 0.12, "状态平面：Checkpoint + Artifact + 上下文压缩 + 长期记忆 + Agent Skill", TEAL, 12)
    _box(ax, 0.04, 0.16, 0.92, 0.12, "治理平面：S-ABAC + 四道 PEP + 审批 + Receipt + 审计", GREEN, 12)
    for y1, y2 in ((0.73, 0.66), (0.54, 0.47), (0.35, 0.28)):
        _arrow(ax, (0.50, y1), (0.50, y2))
    ax.text(0.5, 0.06, "统一 task_id / trace_id / 用户身份 / Agent 与 Tool 契约 / 数据范围", ha="center", fontsize=11, color=_hex(TEXT))
    _save(fig, "overall_architecture.png")


def make_main_agent_console_views() -> None:
    """Compose the five original decision-console screenshots without altering them."""
    screenshots = {
        "intent": SCREENSHOT_DIR / "屏幕截图 2026-08-11 085355.png",
        "entities": SCREENSHOT_DIR / "屏幕截图 2026-08-11 085405.png",
        "subtasks": SCREENSHOT_DIR / "屏幕截图 2026-08-11 085414.png",
        "agents": SCREENSHOT_DIR / "屏幕截图 2026-08-11 085427.png",
        "evidence": SCREENSHOT_DIR / "屏幕截图 2026-08-11 085438.png",
    }
    missing = [str(path) for path in screenshots.values() if not path.exists()]
    if missing:
        raise FileNotFoundError(f"Required decision-console screenshots are missing: {missing}")

    fig = plt.figure(figsize=(12, 14))
    grid = fig.add_gridspec(10, 2, width_ratios=[1.08, 0.92], hspace=0.28, wspace=0.035)
    placements = [
        ("intent", grid[0:3, 0], "（a）意图详情与任务边界"),
        ("entities", grid[3:6, 0], "（b）实体识别与数据范围"),
        ("subtasks", grid[6:9, 0], "（c）子任务拆解与依赖"),
        ("agents", grid[0:6, 1], "（d）Agent 候选与排除原因"),
        ("evidence", grid[6:10, 1], "（e）路由决策依据"),
    ]
    for key, slot, title in placements:
        ax = fig.add_subplot(slot)
        ax.imshow(plt.imread(screenshots[key]))
        ax.set_title(title, fontsize=12, weight="bold", color=_hex(NAVY), pad=7)
        ax.axis("off")
    fig.suptitle("主 Agent 决策台：任务画像、候选过滤与可解释路由", fontsize=18, weight="bold", color=_hex(NAVY), y=0.995)
    _save(fig, "main_agent_console_views.png")


def make_intent_metrics() -> None:
    labels = [
        "主意图准确率",
        "主目标准确率",
        "子意图召回率",
        "实体字段准确率",
        "依赖完全准确率",
        "澄清判定准确率",
    ]
    values = [93.94, 100.00, 98.48, 96.46, 92.31, 95.60]
    colors = [_hex(BLUE), _hex(TEAL), _hex(GREEN), _hex(PURPLE), _hex(ORANGE), _hex(RED)]
    fig, ax = plt.subplots(figsize=(10.5, 5.2))
    bars = ax.barh(labels[::-1], values[::-1], color=colors[::-1], height=0.62)
    ax.set_xlim(0, 105)
    ax.set_xlabel("准确率 / 召回率（%）")
    ax.set_title("离线规则模式：33 条意图评测关键指标", fontsize=16, weight="bold", color=_hex(NAVY))
    ax.grid(axis="x", linestyle="--", alpha=0.25)
    for bar, value in zip(bars, values[::-1]):
        ax.text(value + 0.8, bar.get_y() + bar.get_height() / 2, f"{value:.2f}%", va="center", fontsize=10, color=_hex(TEXT))
    for spine in ("top", "right", "left"):
        ax.spines[spine].set_visible(False)
    _save(fig, "intent_metrics.png")


def make_multi_agent_planes() -> None:
    fig, ax = plt.subplots(figsize=(12, 5.3))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.text(0.5, 0.94, "控制流与业务数据流分离", ha="center", fontsize=18, weight="bold", color=_hex(NAVY))
    ax.text(0.04, 0.72, "控制面", fontsize=13, weight="bold", color=_hex(BLUE))
    xs = [0.16, 0.40, 0.64, 0.82]
    labels = ["TaskGraph\n依赖与状态", "Scheduler\n并行与汇合", "Agent Executor\n受控执行", "StepResult\n标准终态"]
    for x, label in zip(xs, labels):
        _box(ax, x, 0.64, 0.15, 0.16, label, BLUE if x < 0.6 else PURPLE, 10)
    for a, b in zip(xs[:-1], xs[1:]):
        _arrow(ax, (a + 0.15, 0.72), (b, 0.72))
    ax.text(0.04, 0.34, "数据面", fontsize=13, weight="bold", color=_hex(GREEN))
    labels = ["Agent 输出", "Artifact\n所有权/血缘", "ArtifactRef\n脱敏引用", "Resolver + Guard\n受控读取"]
    for x, label in zip(xs, labels):
        _box(ax, x, 0.24, 0.15, 0.16, label, GREEN if x < 0.6 else TEAL, 10)
    for a, b in zip(xs[:-1], xs[1:]):
        _arrow(ax, (a + 0.15, 0.32), (b, 0.32))
    _arrow(ax, (0.715, 0.64), (0.715, 0.40), ORANGE)
    ax.text(0.73, 0.51, "StepResult.outputs", va="center", fontsize=9, color=_hex(ORANGE))
    ax.text(0.5, 0.08, "TaskGraph 描述执行依赖；Artifact 描述数据依赖", ha="center", fontsize=11, weight="bold", color=_hex(TEXT))
    _save(fig, "multi_agent_planes.png")


def make_tool_funnel() -> None:
    fig, ax = plt.subplots(figsize=(10, 4.8))
    values = [48, 2.17]
    labels = ["过滤前\nOffice 30 + Excel 18", "过滤后\n平均候选"]
    colors = [_hex(BLUE), _hex(GREEN)]
    bars = ax.barh(labels, values, color=colors, height=0.55)
    ax.set_xlim(0, 52)
    ax.set_xlabel("候选工具数量")
    ax.set_title("Tool Resolver 候选空间压缩（固定评估集）", fontsize=16, weight="bold", color=_hex(NAVY))
    ax.grid(axis="x", linestyle="--", alpha=0.3)
    for bar, value in zip(bars, values):
        ax.text(value + 0.7, bar.get_y() + bar.get_height() / 2, f"{value:g}", va="center", fontsize=12, weight="bold", color=_hex(TEXT))
    ax.text(26, 0.46, "候选减少 95.5%", ha="center", fontsize=14, color=_hex(ORANGE), weight="bold")
    ax.text(26, -0.42, "Top-1/Top-3 = 100%（audit 模式；不代表开放环境泛化）", ha="center", fontsize=10, color=_hex(GRAY))
    for spine in ("top", "right", "left"):
        ax.spines[spine].set_visible(False)
    _save(fig, "tool_funnel.png")


def make_recovery_benchmark() -> None:
    fig, ax = plt.subplots(figsize=(10, 5.3))
    names = ["B0 无恢复", "B1 有界重试", "B2 重试+等价改派"]
    closure = [12.4, 27.2, 40.0]
    recovered = [0.0, 16.9, 31.5]
    x = range(len(names))
    width = 0.34
    bars1 = ax.bar([i - width / 2 for i in x], closure, width, label="任务闭环率", color=_hex(BLUE))
    bars2 = ax.bar([i + width / 2 for i in x], recovered, width, label="首次失败后恢复成功率", color=_hex(GREEN))
    ax.set_xticks(list(x), names)
    ax.set_ylim(0, 46)
    ax.set_ylabel("比例（%）")
    ax.set_title("失败恢复离线对照实验（每种策略 500 次）", fontsize=16, weight="bold", color=_hex(NAVY))
    ax.legend(frameon=False, loc="upper left")
    ax.grid(axis="y", linestyle="--", alpha=0.3)
    for bars in (bars1, bars2):
        for bar in bars:
            ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + 0.8, f"{bar.get_height():.1f}%", ha="center", fontsize=10)
    ax.text(1.0, -7.5, "固定种子 + Stub Router/Executor；重复副作用=0，治理违规=0", ha="center", fontsize=9, color=_hex(GRAY))
    for spine in ("top", "right"):
        ax.spines[spine].set_visible(False)
    _save(fig, "recovery_benchmark.png")


def make_memory_skill_layers() -> None:
    fig, ax = plt.subplots(figsize=(12, 6.3))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.text(0.5, 0.955, "数字员工面临的真实问题与本项目方案", ha="center", fontsize=18, weight="bold", color=_hex(NAVY))
    ax.text(0.24, 0.875, "真实情况", ha="center", fontsize=14, weight="bold", color=_hex(ORANGE))
    ax.text(0.76, 0.875, "采取的方案", ha="center", fontsize=14, weight="bold", color=_hex(GREEN))
    ax.plot([0.04, 0.44], [0.84, 0.84], color=_hex(ORANGE), linewidth=1.5, alpha=0.7)
    ax.plot([0.56, 0.96], [0.84, 0.84], color=_hex(GREEN), linewidth=1.5, alpha=0.7)

    left = [
        ("长对话持续膨胀", "历史消息、工具结果和 Agent 回传持续占用 Token，\n关键约束容易被淹没。"),
        ("跨会话信息遗忘", "语言、报告风格、文档格式、审批约束和稳定事实\n需要反复说明。"),
        ("重复步骤每次重学", "同一 Agent 反复执行相同工具链，仍需重新探索\n参数与输出契约。"),
    ]
    right = [
        ("分层上下文压缩", "结构化摘要 + 最近两轮；执行中的 Plan、TaskGraph\n和 Artifact 独立保留。"),
        ("受治理的长期记忆", "固定办公标签、后台提取、来源轮次、冲突替换、\n用户/项目隔离和时间衰减。"),
        ("步骤 / 智能体 Skill", "LLM 反思 + 两次独立成功证据 + 契约校验，\n仅绑定对应 Plan 步骤。"),
    ]

    def card(x: float, y: float, title: str, body: str, color: str, fill: str) -> None:
        patch = FancyBboxPatch(
            (x, y), 0.40, 0.18,
            boxstyle="round,pad=0.012,rounding_size=0.018",
            linewidth=1.3, edgecolor=_hex(color), facecolor=_hex(fill),
        )
        ax.add_patch(patch)
        ax.text(x + 0.025, y + 0.125, title, ha="left", va="center", fontsize=12, weight="bold", color=_hex(color))
        ax.text(x + 0.025, y + 0.065, body, ha="left", va="center", fontsize=9.2, color=_hex(TEXT), linespacing=1.35)

    for y, problem, solution in zip((0.60, 0.36, 0.12), left, right):
        card(0.04, y, problem[0], problem[1], ORANGE, "FCF3E8")
        card(0.56, y, solution[0], solution[1], GREEN, "EDF7F3")
        _arrow(ax, (0.45, y + 0.09), (0.55, y + 0.09))
    _save(fig, "memory_skill_layers.png")


def make_context_tokens() -> None:
    fig, ax = plt.subplots(figsize=(10, 5.2))
    names = ["8K", "16K", "32K"]
    before = [4426, 6939, 11905]
    after = [3460, 6262, 6137]
    x = range(len(names))
    width = 0.36
    bars1 = ax.bar([i - width / 2 for i in x], before, width, label="压缩前 Token", color=_hex(BLUE), alpha=0.86)
    bars2 = ax.bar([i + width / 2 for i in x], after, width, label="压缩后 Token", color=_hex(TEAL), alpha=0.92)
    ax.set_xticks(list(x), names)
    ax.set_ylabel("估算 Token 数")
    ax.set_title("真实多轮上下文压缩：8K / 16K / 32K", fontsize=16, weight="bold", color=_hex(NAVY))
    ax.legend(frameon=False, loc="upper left")
    ax.grid(axis="y", linestyle="--", alpha=0.25)
    for bars in (bars1, bars2):
        for bar in bars:
            ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + 140, f"{int(bar.get_height()):,}", ha="center", fontsize=9)
    for spine in ("top", "right"):
        ax.spines[spine].set_visible(False)
    _save(fig, "context_tokens.png")


def make_context_continuity() -> None:
    fig, ax = plt.subplots(figsize=(10, 5.2))
    labels = ["报告核心事实", "结构化产物", "执行步骤", "接续工作流"]
    values = [9, 9, 9, 3]
    totals = [9, 9, 9, 3]
    bars = ax.barh(labels, values, color=_hex(TEAL), alpha=0.92)
    ax.invert_yaxis()
    ax.set_xlim(0, 10)
    ax.set_xlabel("压缩后实际完成数量")
    ax.set_title("压缩后任务接续：从历史恢复到完整工作流执行", fontsize=16, weight="bold", color=_hex(NAVY))
    ax.grid(axis="x", linestyle="--", alpha=0.22)
    for bar, value, total in zip(bars, values, totals):
        ax.text(value + 0.12, bar.get_y() + bar.get_height() / 2, f"{value}/{total}", va="center", fontsize=11, weight="bold")
    ax.text(9.85, 3.35, "禁止的外部副作用：0 次", ha="right", fontsize=10, color=_hex(GREEN), bbox={"boxstyle": "round,pad=0.35", "fc": "white", "ec": _hex(TEAL)})
    for spine in ("top", "right", "left"):
        ax.spines[spine].set_visible(False)
    _save(fig, "context_continuity.png")


def make_memory_decay() -> None:
    fig, ax = plt.subplots(figsize=(9.5, 5))
    days = [0, 30, 90, 180]
    weights = [0.855, 0.679, 0.428, 0.214]
    ax.plot(days, weights, marker="o", linewidth=2.6, markersize=8, color=_hex(TEAL))
    ax.fill_between(days, weights, alpha=0.12, color=_hex(TEAL))
    for d, w in zip(days, weights):
        ax.annotate(f"{w:.3f}", (d, w), textcoords="offset points", xytext=(0, 10), ha="center", fontsize=10)
    ax.set_title("长期记忆时间衰减（来源事实不删除）", fontsize=16, weight="bold", color=_hex(NAVY))
    ax.set_xlabel("距上次强化天数")
    ax.set_ylabel("召回权重")
    ax.set_ylim(0, 1.0)
    ax.grid(linestyle="--", alpha=0.3)
    ax.text(90, 0.08, "Memory 治理专项：64/64\nPrecision = Recall = F1 = 1.0\n（确定性模块评测）", ha="center", fontsize=10, color=_hex(TEXT), bbox={"boxstyle": "round,pad=0.4", "fc": _hex(LIGHT), "ec": _hex(MID)})
    for spine in ("top", "right"):
        ax.spines[spine].set_visible(False)
    _save(fig, "memory_decay.png")


def make_skill_results() -> None:
    labels = ["首次证据建档", "独立证据晋升", "实体变化复用", "契约漂移拒绝", "用户/范围隔离", "反思/身份门禁"]
    values = [6, 6, 6, 6, 6, 6]
    colors = [_hex(TEAL), _hex(BLUE), _hex(GREEN), _hex(ORANGE), _hex(PURPLE), _hex(RED)]
    fig, ax = plt.subplots(figsize=(10, 5.5))
    bars = ax.barh(labels[::-1], values[::-1], color=colors[::-1])
    ax.set_xlim(0, 7.2)
    ax.set_xlabel("符合预期案例数（每类 6 条）")
    ax.set_title("36 例 Skill 能力判定结果（六个维度均为 6/6）", fontsize=16, weight="bold", color=_hex(NAVY))
    ax.grid(axis="x", linestyle="--", alpha=0.25)
    for bar, value in zip(bars, values[::-1]):
        ax.text(value + 0.1, bar.get_y() + bar.get_height() / 2, f"{value}/6", va="center", fontsize=10, weight="bold")
    for spine in ("top", "right", "left"):
        ax.spines[spine].set_visible(False)
    _save(fig, "skill_results.png")


def prepare_assets() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    make_overall_architecture()
    make_main_agent_console_views()
    make_intent_metrics()
    make_multi_agent_planes()
    make_tool_funnel()
    make_recovery_benchmark()
    make_memory_skill_layers()
    make_context_tokens()
    make_context_continuity()
    make_memory_decay()
    make_skill_results()

    copies = {
        "image2.png": "main_agent_decision_flow.png",
        "image6.png": "main_agent_dashboard.png",
        "image3.png": "security_gates.png",
        "image10.png": "permission_dashboard.png",
        "image5.png": "governance_validation.png",
    }
    for source_name, target_name in copies.items():
        source = EXTRACTED_MEDIA / source_name
        if not source.exists():
            raise FileNotFoundError(f"Required report asset is missing: {source}")
        shutil.copy2(source, ASSET_DIR / target_name)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 70, start: int = 80, bottom: int = 70, end: int = 80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size: float = 12, bold: bool | None = None, color: str = TEXT, latin: str = "Times New Roman", east_asia: str = "宋体") -> None:
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.25)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.25)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for level in (1, 2, 3):
        style = styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(14)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.paragraph_format.space_before = Pt(12 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    if "Figure Caption" not in styles:
        caption = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
        caption.font.name = "Times New Roman"
        caption.font.size = Pt(9)
        caption._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(6)

    if "Table Caption" not in styles:
        caption = styles.add_style("Table Caption", WD_STYLE_TYPE.PARAGRAPH)
        caption.font.name = "Times New Roman"
        caption.font.size = Pt(9)
        caption.font.bold = True
        caption._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(5)
        caption.paragraph_format.space_after = Pt(3)
        caption.paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Consolas"
        code.font.size = Pt(8.5)
        code._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        code.paragraph_format.left_indent = Cm(0.55)
        code.paragraph_format.right_indent = Cm(0.25)
        code.paragraph_format.space_before = Pt(3)
        code.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("面向数字员工的智能任务执行关键技术研究")
    set_run_font(run, 8.5, color=GRAY, east_asia="微软雅黑")
    p_pr = header._p.get_or_add_pPr()
    bottom = OxmlElement("w:pBdr")
    p_pr.append(bottom)
    border = OxmlElement("w:bottom")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), "4")
    border.set(qn("w:color"), MID)
    bottom.append(border)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("—  ")
    set_run_font(run, 9, color=GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    run2 = footer.add_run("  —")
    set_run_font(run2, 9, color=GRAY)


def add_cover(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("面向数字员工的\n智能任务执行关键技术研究")
    set_run_font(run, 22, True, NAVY, east_asia="微软雅黑")
    p.paragraph_format.space_after = Pt(16)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p2.add_run("——主 Agent 决策、多 Agent 编排、记忆与 Skill、权限治理")
    set_run_font(run, 13.5, True, BLUE, east_asia="微软雅黑")
    p2.paragraph_format.space_after = Pt(32)

    table = doc.add_table(rows=4, cols=2)
    table.alignment = 1
    table.autofit = False
    values = [("报告类型", "实习课题技术报告"), ("课题名称", "面向数字员工的智能任务执行关键技术研究"), ("编写单位", "数字员工智能任务执行关键技术研究组"), ("日期", "2026 年 8 月")]
    for row, (key, value) in zip(table.rows, values):
        row.cells[0].width = Cm(3.2)
        row.cells[1].width = Cm(10.5)
        for cell in row.cells:
            set_cell_margins(cell, 100, 120, 100, 120)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(row.cells[0], LIGHT)
        r1 = row.cells[0].paragraphs[0].add_run(key)
        set_run_font(r1, 10.5, True, NAVY, east_asia="微软雅黑")
        r2 = row.cells[1].paragraphs[0].add_run(value)
        set_run_font(r2, 10.5)
    doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = note.add_run("A4 单栏学术版式 · IEEE 风格顺序编码引用")
    set_run_font(r, 9, color=GRAY, east_asia="微软雅黑")
    doc.add_page_break()


def add_toc(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("目  录")
    set_run_font(r, 16, True, NAVY, east_asia="微软雅黑")
    toc_p = doc.add_paragraph()
    run = toc_p.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "打开 Word 后右键此处并选择‘更新域’，即可生成目录。"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, text, end])
    doc.add_page_break()


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def add_inline(paragraph, text: str, size: float = 12) -> None:
    for part in INLINE_RE.split(text):
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        code = part.startswith("`") and part.endswith("`")
        value = part[2:-2] if bold else part[1:-1] if code else part
        run = paragraph.add_run(value)
        # Inline technical terms remain part of the body text and therefore use
        # the same Times New Roman/宋体 pairing and 小四 size as surrounding prose.
        set_run_font(run, size, bold if bold else None, TEXT, latin="Times New Roman", east_asia="宋体")


def add_body_paragraph(doc: Document, text: str, indent: bool = True) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.74) if indent else None
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(4)
    add_inline(p, text)


def add_list_item(doc: Document, text: str, numbered: bool) -> None:
    style = "List Number" if numbered else "List Bullet"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.3)
    p.paragraph_format.space_after = Pt(2)
    add_inline(p, text)


def add_code_block(doc: Document, lines: list[str]) -> None:
    p = doc.add_paragraph(style="Code Block")
    set_cell_like_background(p, "F1F4F7")
    for index, line in enumerate(lines):
        if index:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, 8.5, color=TEXT, latin="Consolas", east_asia="微软雅黑")


def set_cell_like_background(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_equation(doc: Document, lines: list[str]) -> None:
    text = " ".join(line.strip() for line in lines if line.strip())
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, 10.5, color=NAVY, latin="Cambria Math", east_asia="宋体")


def add_markdown_table(doc: Document, rows: list[list[str]], caption_text: str) -> None:
    if len(rows) < 2:
        return
    rows = [row for row in rows if not all(re.fullmatch(r"\s*:?-+:?\s*", cell) for cell in row)]
    if not rows:
        return
    caption = doc.add_paragraph(style="Table Caption")
    add_inline(caption, caption_text, 9)
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.alignment = 1
    table.autofit = True
    font_size = 8.2 if columns >= 5 else 8.8 if columns == 4 else 9.2
    for r_idx, values in enumerate(rows):
        row = table.rows[r_idx]
        if r_idx == 0:
            set_repeat_table_header(row)
        for c_idx in range(columns):
            cell = row.cells[c_idx]
            value = values[c_idx].strip() if c_idx < len(values) else ""
            cell.text = ""
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_shading(cell, NAVY)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, LIGHT)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 or re.fullmatch(r"[\d.,%/（）()～~-]+", value) else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, value, font_size)
            for run in p.runs:
                if r_idx == 0:
                    set_run_font(
                        run,
                        font_size,
                        True,
                        "FFFFFF",
                        east_asia="微软雅黑",
                    )
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_figure(doc: Document, alt: str, path_text: str) -> None:
    path = (SOURCE.parent / path_text).resolve()
    if not path.exists():
        raise FileNotFoundError(f"Figure not found: {path}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(15.8))
    p.paragraph_format.space_after = Pt(2)
    caption = doc.add_paragraph(style="Figure Caption")
    add_inline(caption, alt, 9)


def table_caption(rows: list[list[str]], number: int) -> str:
    headers = tuple(cell.strip() for cell in rows[0]) if rows else ()
    titles = {
        ("模块", "主要研究内容", "主要成果证据"): "研究分工与贡献边界",
        ("方案/产品", "决策方式", "优势", "对本课题的不足", "本项目取舍"): "主 Agent 同类产品与技术路线对比",
        ("指标", "结果", "解释"): "主 Agent 意图画像评测结果",
        ("方案/产品", "主要机制", "优势", "局限", "本项目吸收点"): "多 Agent 编排方案对比",
        ("方案", "解决重点", "对本课题的启示", "不直接采用的原因"): "工具选择技术路线对比",
        ("指标", "结果", "测试条件"): "Tool Resolver 固定评估结果",
        ("策略", "任务闭环率", "首次失败后恢复成功率", "重复副作用", "治理违规"): "失败恢复离线对照实验",
        ("产品/研究", "核心机制", "优势", "对银行办公场景的不足", "本项目对应设计"): "记忆与 Skill 同类产品对比",
        ("上下文配置", "压缩代数", "压缩前 Token", "压缩后 Token", "降幅"): "上下文压缩评测结果",
        ("方案/产品", "主要依据", "优势", "局限", "本项目定位"): "权限模型与产品对比",
        ("测试组", "结果", "主要验证点"): "权限治理测试结果",
        ("能力", "主要实现或证据"): "关键实现与证据索引",
    }
    title = titles.get(headers, "技术方案与评测结果")
    return f"表 {number}  {title}"


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_heading(text, level=level)
    # Apply the requested heading typography directly so Word keeps it after
    # field/TOC updates, even when the document theme rewrites style fonts.
    for run in p.runs:
        set_run_font(run, 14, bold=True, color="000000", east_asia="黑体")
    if level == 1 and re.match(r"^[1-4] ", text):
        p.paragraph_format.page_break_before = True


def parse_markdown(doc: Document) -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    table_number = 0
    # Skip the cover title, subtitle and metadata. Content starts at the research background.
    while index < len(lines) and lines[index].strip() != "## 研究背景":
        index += 1
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if not line:
            index += 1
            continue
        if line.startswith("```"):
            block: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            add_code_block(doc, block)
            index += 1
            continue
        if line == "$$":
            block = []
            index += 1
            while index < len(lines) and lines[index].strip() != "$$":
                block.append(lines[index])
                index += 1
            add_equation(doc, block)
            index += 1
            continue
        image_match = re.fullmatch(r"!\[(.+?)\]\((.+?)\)", line)
        if image_match:
            add_figure(doc, image_match.group(1), image_match.group(2))
            index += 1
            continue
        if line.startswith("|"):
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = [[cell.strip() for cell in row.strip("|").split("|")] for row in table_lines]
            table_number += 1
            add_markdown_table(doc, rows, table_caption(rows, table_number))
            continue
        if line.startswith("# "):
            add_heading(doc, line[2:].strip(), 1)
            index += 1
            continue
        if line.startswith("## "):
            text = line[3:].strip()
            if text in {"研究背景", "研究目标与总体技术框架", "总体结论"} or text.startswith("附录"):
                add_heading(doc, text, 1)
            elif text.startswith("——"):
                index += 1
                continue
            else:
                add_heading(doc, text, 2)
            index += 1
            continue
        if line.startswith("### "):
            add_heading(doc, line[4:].strip(), 3)
            index += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            add_list_item(doc, re.sub(r"^\d+\.\s+", "", line), numbered=True)
            index += 1
            continue
        if line.startswith("- "):
            add_list_item(doc, line[2:].strip(), numbered=False)
            index += 1
            continue
        if line.startswith("**图") and line.endswith("**"):
            p = doc.add_paragraph(style="Figure Caption")
            add_inline(p, line[2:-2], 9)
            index += 1
            continue
        add_body_paragraph(doc, line, indent=not line.startswith("["))
        index += 1


def build_report() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    prepare_assets()
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "面向数字员工的智能任务执行关键技术研究"
    doc.core_properties.subject = "主 Agent 决策、多 Agent 编排、记忆与 Skill、权限治理"
    doc.core_properties.author = "数字员工智能任务执行关键技术研究组"
    doc.core_properties.keywords = "数字员工, 主Agent, 多智能体, MCP, 长期记忆, Agent Skill, S-ABAC"
    add_cover(doc)
    add_toc(doc)
    parse_markdown(doc)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_report())
