from __future__ import annotations

import math
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import (
    WD_ALIGN_VERTICAL,
    WD_CELL_VERTICAL_ALIGNMENT,
    WD_ROW_HEIGHT_RULE,
    WD_TABLE_ALIGNMENT,
)
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "output"
ASSET_DIR = OUTPUT_DIR / "main_agent_governance_report_assets"
OUTPUT_PATH = OUTPUT_DIR / "面向数字员工的主Agent意图识别路由与权限治理研究报告.docx"

NAVY = "163A59"
BLUE = "256FA1"
TEAL = "1B9AAA"
PALE_BLUE = "EAF3F8"
PALE_TEAL = "E8F6F6"
PALE_GOLD = "FFF5DC"
PALE_RED = "FCECEC"
MID_GREY = "D9E2E8"
LIGHT_GREY = "F5F7F9"
DARK = "253746"
WHITE = "FFFFFF"


def _font_name() -> str:
    candidates = ["Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "Arial Unicode MS"]
    available = {item.name for item in font_manager.fontManager.ttflist}
    return next((name for name in candidates if name in available), "DejaVu Sans")


CHART_FONT = _font_name()
plt.rcParams["font.sans-serif"] = [CHART_FONT]
plt.rcParams["axes.unicode_minus"] = False


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key in ("val", "sz", "space", "color"):
            if key in kwargs[edge]:
                element.set(qn("w:" + key), str(kwargs[edge][key]))


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="宋体", size=10.5, bold=None, color=None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_repeat_on_open(doc: Document) -> None:
    settings = doc.settings._element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr_text, fld_char2])
    set_run_font(run, "微软雅黑", 9, color="647987")


def add_toc(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "打开 Word 后右键此处并选择“更新域”，即可生成目录。"
    separate.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])
    set_run_font(run, "宋体", 10.5, color="647987")


def add_para(doc: Document, text: str = "", *, bold_prefix: str | None = None,
             align=None, indent=True, space_after=5, style=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(space_after)
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, "宋体", 10.5, True, DARK)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest, "宋体", 10.5, False, DARK)
    else:
        run = p.add_run(text)
        set_run_font(run, "宋体", 10.5, False, DARK)
    return p


def add_bullets(doc: Document, items: list[str], level=0) -> None:
    for text in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.left_indent = Cm(0.75 + level * 0.55)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text)
        set_run_font(run, "宋体", 10.5, color=DARK)


def add_numbered(doc: Document, items: list[str]) -> None:
    for text in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text)
        set_run_font(run, "宋体", 10.5, color=DARK)


def add_callout(doc: Document, title: str, text: str, fill=PALE_BLUE, accent=BLUE) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(0.18)
    table.columns[1].width = Cm(15.9)
    left, body = table.rows[0].cells
    set_cell_shading(left, accent)
    set_cell_shading(body, fill)
    set_cell_margins(left, 0, 0, 0, 0)
    set_cell_margins(body, 150, 180, 150, 180)
    left.width = Cm(0.18)
    body.width = Cm(15.9)
    p = body.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, "微软雅黑", 10.5, True, accent)
    p2 = body.add_paragraph()
    p2.paragraph_format.line_spacing = 1.35
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(text)
    set_run_font(r2, "宋体", 10, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code(doc: Document, text: str, caption: str | None = None) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F1F4F6")
    set_cell_border(cell, top={"val": "single", "sz": 6, "color": MID_GREY},
                    bottom={"val": "single", "sz": 6, "color": MID_GREY},
                    left={"val": "single", "sz": 6, "color": MID_GREY},
                    right={"val": "single", "sz": 6, "color": MID_GREY})
    set_cell_margins(cell, 130, 180, 130, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.1
    for idx, line in enumerate(text.splitlines()):
        if idx:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, "Consolas", 8.8, color="29434E")
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(6)
        rr = cp.add_run(caption)
        set_run_font(rr, "宋体", 9, color="647987")


def add_table(doc: Document, headers: list[str], rows: list[list[str]], *, widths=None,
              font_size=9.2, header_fill=NAVY, zebra=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    for i, title in enumerate(headers):
        set_cell_shading(hdr[i], header_fill)
        set_cell_margins(hdr[i], 100, 90, 100, 90)
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(title))
        set_run_font(r, "微软雅黑", font_size, True, WHITE)
        if widths:
            hdr[i].width = Cm(widths[i])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        if zebra and ridx % 2 == 1:
            for cell in cells:
                set_cell_shading(cell, LIGHT_GREY)
        for i, value in enumerate(row):
            set_cell_margins(cells[i], 90, 90, 90, 90)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cells[i].width = Cm(widths[i])
            p = cells[i].paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(value)) < 16 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, "宋体", font_size, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    set_run_font(r, "宋体", 9, color="647987")


def add_figure(doc: Document, path: Path, caption: str, width=15.5) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width))
    add_caption(doc, caption)


def add_screenshot_placeholder(doc: Document, title: str, guidance: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.rows[0].height = Cm(6.8)
    table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_shading(cell, "F3F6F8")
    set_cell_border(cell, top={"val": "dashed", "sz": 12, "color": "9EB2BF"},
                    bottom={"val": "dashed", "sz": 12, "color": "9EB2BF"},
                    left={"val": "dashed", "sz": 12, "color": "9EB2BF"},
                    right={"val": "dashed", "sz": 12, "color": "9EB2BF"})
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("前端截图待粘贴")
    set_run_font(r, "微软雅黑", 16, True, "7B919F")
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(guidance)
    set_run_font(r2, "宋体", 9.5, color="7B919F")
    add_caption(doc, title)


def _box(ax, xy, w, h, text, color, *, fontsize=10, text_color="white"):
    patch = FancyBboxPatch(xy, w, h, boxstyle="round,pad=0.02,rounding_size=0.025",
                           linewidth=1.2, edgecolor=color, facecolor=color)
    ax.add_patch(patch)
    ax.text(xy[0] + w / 2, xy[1] + h / 2, text, ha="center", va="center",
            fontsize=fontsize, color=text_color, fontweight="bold")
    return patch


def _arrow(ax, start, end, color="#6F8795"):
    ax.add_patch(FancyArrowPatch(start, end, arrowstyle="-|>", mutation_scale=12,
                                color=color, linewidth=1.4))


def make_architecture_figure(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(13.5, 6.2), dpi=180)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    y = 0.72
    boxes = [
        (0.02, "用户请求\n+身份/会话", "#3B5B73"),
        (0.16, "Task Profiler\n规则+语义融合", "#256FA1"),
        (0.31, "TaskProfile\n结构化任务画像", "#1B9AAA"),
        (0.46, "权限硬过滤\nAuthorized IDs", "#C9782B"),
        (0.61, "AgentCard 路由\n能力匹配+打分", "#6D5AA7"),
        (0.76, "RoutingDecision\n派发/澄清/拒绝", "#31785E"),
        (0.90, "受控执行\nAgent / Tool", "#3B5B73"),
    ]
    widths = [0.105, 0.115, 0.115, 0.115, 0.115, 0.115, 0.085]
    for i, (x, text, color) in enumerate(boxes):
        _box(ax, (x, y), widths[i], 0.15, text, color, fontsize=8.2)
        if i < len(boxes) - 1:
            _arrow(ax, (x + widths[i], y + 0.075), (boxes[i + 1][0] - 0.008, y + 0.075))
    _box(ax, (0.21, 0.30), 0.17, 0.14, "S-ABAC PDP\nSubject/Object/Scenario/Action", "#B34E4E", fontsize=8.5)
    _box(ax, (0.43, 0.30), 0.14, 0.14, "人工审批\n一次性签名绑定", "#C9782B", fontsize=8.5)
    _box(ax, (0.62, 0.30), 0.14, 0.14, "Artifact Guard\n所有权/敏感级别", "#1B9AAA", fontsize=8.5)
    _box(ax, (0.81, 0.30), 0.14, 0.14, "Receipt / 审计\n防重与可追溯", "#31785E", fontsize=8.5)
    for x in (0.295, 0.50, 0.69, 0.88):
        _arrow(ax, (x, 0.44), (x, 0.69))
    ax.text(0.02, 0.10, "核心原则：权限是硬约束，不参与软打分；模型负责理解，确定性组件负责边界、校验和执行。",
            fontsize=11, color="#253746", fontweight="bold")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def make_decision_flow(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(11.8, 7.0), dpi=180)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    nodes = [
        ((0.38, 0.86), 0.24, 0.09, "自然语言请求", "#3B5B73"),
        ((0.31, 0.70), 0.38, 0.10, "意图目录 + 规则识别 + 可选语义识别", "#256FA1"),
        ((0.31, 0.54), 0.38, 0.10, "融合、去重、否定识别、依赖补全", "#1B9AAA"),
        ((0.35, 0.38), 0.30, 0.10, "TaskProfile / Pydantic 校验", "#6D5AA7"),
    ]
    for xy, w, h, text, color in nodes:
        _box(ax, xy, w, h, text, color, fontsize=9.5)
    for y1, y2 in ((0.86, 0.80), (0.70, 0.64), (0.54, 0.48)):
        _arrow(ax, (0.50, y1), (0.50, y2))
    _box(ax, (0.35, 0.22), 0.30, 0.10, "关键字段是否完整？", "#C9782B", fontsize=10)
    _arrow(ax, (0.50, 0.38), (0.50, 0.32))
    _box(ax, (0.05, 0.05), 0.25, 0.09, "CLARIFY\n生成具体澄清问题", "#B34E4E", fontsize=9.5)
    _box(ax, (0.38, 0.05), 0.25, 0.09, "权限过滤 + 路由打分", "#31785E", fontsize=9.5)
    _box(ax, (0.70, 0.05), 0.25, 0.09, "REJECT\n无合法/有能力候选", "#737F88", fontsize=9.5)
    _arrow(ax, (0.35, 0.27), (0.18, 0.14))
    _arrow(ax, (0.50, 0.22), (0.50, 0.14))
    _arrow(ax, (0.65, 0.27), (0.82, 0.14))
    ax.text(0.25, 0.21, "缺失", fontsize=8.5, color="#B34E4E")
    ax.text(0.52, 0.17, "完整", fontsize=8.5, color="#31785E")
    ax.text(0.72, 0.21, "无候选", fontsize=8.5, color="#737F88")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def make_security_gates(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(12.5, 5.6), dpi=180)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    gates = [
        (0.03, "闸门一\n路由前", "过滤不可见 Agent\n限定合法候选集", "#256FA1"),
        (0.27, "闸门二\n派发前", "校验动作、数据域\n风险与场景", "#6D5AA7"),
        (0.51, "闸门三\n工具前", "逐次鉴权\n审批与参数绑定", "#C9782B"),
        (0.75, "闸门四\n结果前", "Artifact 所有权\n敏感级别与审计", "#1B9AAA"),
    ]
    for i, (x, title, body, color) in enumerate(gates):
        _box(ax, (x, 0.48), 0.19, 0.25, title, color, fontsize=12)
        ax.text(x + 0.095, 0.35, body, ha="center", va="center", fontsize=9.5, color="#253746")
        if i < len(gates) - 1:
            _arrow(ax, (x + 0.19, 0.605), (gates[i + 1][0] - 0.015, 0.605))
    ax.text(0.5, 0.12, "任一关键属性缺失、资源未知、校验失败或存储损坏 → Fail Closed（拒绝或转人工核对）",
            ha="center", fontsize=10.5, fontweight="bold", color="#B34E4E")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def make_metrics_chart(path: Path) -> None:
    labels = ["主意图准确率", "主目标准确率", "子意图召回率", "实体字段准确率", "依赖完全准确率", "澄清判定准确率"]
    values = [93.94, 100.0, 98.48, 96.46, 92.31, 75.0]
    colors = ["#256FA1", "#1B9AAA", "#31785E", "#6D5AA7", "#C9782B", "#B34E4E"]
    fig, ax = plt.subplots(figsize=(11.5, 5.6), dpi=180)
    bars = ax.barh(labels[::-1], values[::-1], color=colors[::-1], height=0.62)
    ax.set_xlim(0, 105)
    ax.set_xlabel("准确率 / 召回率（%）")
    ax.grid(axis="x", linestyle="--", alpha=0.25)
    ax.spines[["top", "right", "left"]].set_visible(False)
    for bar, value in zip(bars, values[::-1]):
        ax.text(value + 1.0, bar.get_y() + bar.get_height() / 2, f"{value:.2f}%",
                va="center", fontsize=9, color="#253746")
    ax.set_title("离线规则模式：33 条意图评测关键指标", fontsize=13, fontweight="bold", color="#163A59")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def make_validation_chart(path: Path) -> None:
    groups = ["主 Agent\n结构/合同测试", "权限治理\n核心测试", "权限 API\n决策矩阵", "专题回归\n含 UI 断言"]
    passed = [113, 105, 12, 252]
    total = [113, 105, 12, 253]
    rates = [p / t * 100 for p, t in zip(passed, total)]
    fig, ax = plt.subplots(figsize=(10.8, 5.4), dpi=180)
    bars = ax.bar(groups, rates, color=["#256FA1", "#1B9AAA", "#31785E", "#C9782B"], width=0.55)
    ax.set_ylim(0, 105)
    ax.set_ylabel("通过率（%）")
    ax.grid(axis="y", linestyle="--", alpha=0.25)
    ax.spines[["top", "right"]].set_visible(False)
    for bar, p, t, rate in zip(bars, passed, total, rates):
        ax.text(bar.get_x() + bar.get_width() / 2, rate + 1.2, f"{p}/{t}\n{rate:.1f}%",
                ha="center", va="bottom", fontsize=9, color="#253746")
    ax.set_title("原型验证结果（2026-08-10 本地复现）", fontsize=13, fontweight="bold", color="#163A59")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.25)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.45)
    section.right_margin = Cm(2.25)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(5)

    for name, size, color, before, after in (
        ("Title", 25, NAVY, 0, 10),
        ("Heading 1", 16, NAVY, 18, 9),
        ("Heading 2", 13, BLUE, 13, 6),
        ("Heading 3", 11, TEAL, 9, 4),
    ):
        style = styles[name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("SUPERAGENT  |  主 Agent 决策与权限治理研究报告")
    set_run_font(r, "微软雅黑", 8.5, color="647987")
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), MID_GREY)
    borders.append(bottom)
    p_pr.append(borders)

    footer = section.footer
    add_page_number(footer.paragraphs[0])
    set_repeat_on_open(doc)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(42)
    r = p.add_run("SUPERAGENT  /  RESEARCH REPORT")
    set_run_font(r, "微软雅黑", 11, True, TEAL)

    band = doc.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = band.cell(0, 0)
    set_cell_shading(cell, NAVY)
    set_cell_margins(cell, 460, 300, 460, 300)
    p1 = cell.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p1.paragraph_format.space_after = Pt(10)
    r1 = p1.add_run("面向数字员工的智能任务执行关键技术研究")
    set_run_font(r1, "微软雅黑", 14, True, "B8DFEA")
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p2.paragraph_format.space_after = Pt(10)
    r2 = p2.add_run("主 Agent 意图识别、权限感知路由\n与权限治理研究报告")
    set_run_font(r2, "微软雅黑", 25, True, WHITE)
    p3 = cell.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p3.paragraph_format.space_after = Pt(0)
    r3 = p3.add_run("—— 原型方案调研、系统设计与量化评测")
    set_run_font(r3, "微软雅黑", 12, False, "DCEAF1")

    doc.add_paragraph().paragraph_format.space_after = Pt(38)
    info = doc.add_table(rows=4, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.RIGHT
    info.autofit = False
    labels = ["报告类型", "研究范围", "实习生", "日期"]
    values = ["课题专题研究报告", "主 Agent 决策与权限治理", "刘建杰", "2026 年 8 月"]
    for i, (label, value) in enumerate(zip(labels, values)):
        c1, c2 = info.rows[i].cells
        c1.width = Cm(3.0)
        c2.width = Cm(7.0)
        set_cell_margins(c1, 80, 120, 80, 120)
        set_cell_margins(c2, 80, 120, 80, 120)
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rr = p1.add_run(label)
        set_run_font(rr, "微软雅黑", 10, True, "647987")
        rr2 = c2.paragraphs[0].add_run(value)
        set_run_font(rr2, "宋体", 10.5, color=DARK)
    doc.add_page_break()


def build_report() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    figures = {
        "architecture": ASSET_DIR / "architecture.png",
        "decision": ASSET_DIR / "decision_flow.png",
        "security": ASSET_DIR / "security_gates.png",
        "metrics": ASSET_DIR / "intent_metrics.png",
        "validation": ASSET_DIR / "validation_results.png",
    }
    make_architecture_figure(figures["architecture"])
    make_decision_flow(figures["decision"])
    make_security_gates(figures["security"])
    make_metrics_chart(figures["metrics"])
    make_validation_chart(figures["validation"])

    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "面向数字员工的主Agent意图识别路由与权限治理研究报告"
    doc.core_properties.subject = "主 Agent 决策、意图识别、权限感知路由与 S-ABAC 权限治理"
    doc.core_properties.author = "刘建杰"
    doc.core_properties.keywords = "数字员工, 主Agent, 意图识别, Agent路由, S-ABAC, 权限治理"
    add_cover(doc)

    doc.add_heading("摘  要", level=1)
    add_para(doc, "随着数字员工从单轮问答向跨系统、跨步骤和带业务副作用的任务执行演进，系统面临两个直接决定可用性的关键问题：第一，主 Agent 能否把自然语言请求稳定转换为可验证的业务意图、实体、动作、依赖和风险画像，并把任务路由到真正有能力且有权限的执行 Agent；第二，系统能否在模型输出不稳定、工具边界复杂、任务可能中断恢复的条件下，保证越权操作被阻断、高风险动作经过审批、敏感数据受控流转且副作用不会重复发生。")
    add_para(doc, "本报告围绕上述问题，对纯规则、纯大模型、向量召回、Planner 直接选 Agent、去中心化 Agent handoff 和“结构化任务画像 + 能力卡 + 权限硬过滤 + 确定性打分”等主 Agent 决策方案进行比较；同时对 ACL、RBAC、ABAC、ReBAC、XACML、OPA 和 Cedar 等权限方案进行调研。在原型规模、可解释性、离线可复现和后续可迁移性的约束下，最终采用规则与可选语义模型融合的 TaskProfile、AgentCard 驱动的权限感知路由，以及 Subject–Object–Scenario–Action 四维 S-ABAC 权限模型。")
    add_para(doc, "原型形成了路由前、派发前、工具调用前和结果返回前四道治理闸门，并将 ALLOW、DENY、REVIEW_REQUIRED 三态决策与一次性审批、Artifact 所有权守卫、审计、幂等键和持久化 Receipt 结合。2026 年 8 月 10 日的本地复现结果显示：33 条离线规则模式意图评测中，主意图准确率为 93.94%，子意图召回率为 98.48%，实体字段准确率为 96.46%，依赖完全准确率为 92.31%，否定动作误执行率为 0；主 Agent 决策相关结构与合同测试 113/113 通过，权限治理核心测试 105/105 通过，权限 API 矩阵 12/12 通过。严格整例通过率为 60.61%，主要短板集中在置信度区间、模糊请求澄清和少数同义表达边界，报告据此提出了后续语义融合评测、阈值校准与外置策略引擎迁移方向。")
    add_callout(doc, "研究结论", "模型适合负责语义理解和候选判断，但权限边界、Schema 校验、审批绑定和副作用防重必须由确定性组件执行。权限不应作为路由分数的一项权重，而应在模型选择之前形成不可绕过的合法候选集合。", PALE_TEAL, TEAL)
    add_para(doc, "关键词：数字员工；主 Agent；意图识别；任务画像；Agent 路由；S-ABAC；人工审批；Artifact；幂等执行", indent=False)

    doc.add_page_break()
    doc.add_heading("目  录", level=1)
    add_toc(doc)
    doc.add_page_break()

    doc.add_heading("1 研究背景与课题范围", level=1)
    doc.add_heading("1.1 研究背景", level=2)
    add_para(doc, "任务书指出，数字员工体系已在流程自动化和智能辅助方面形成基础，但复杂办公任务仍存在任务理解不准、多智能体协同不足和执行过程缺乏有效治理等问题[1]。人员查询、日程管理、课程检索、员工差旅、会议助手和消息发送等场景，看似是自然语言交互，实际往往同时包含对象识别、时间与人员实体抽取、读写动作判定、数据依赖、跨部门能力选择和敏感数据控制。若主 Agent 对其中任一环节判断错误，后续工具调用即使技术上成功，也可能得到业务上错误甚至越权的结果。")
    add_para(doc, "传统问答系统的成功标准是“回答看起来合理”，数字员工的成功标准则是“执行对象正确、执行顺序正确、权限合法、结果可验证、失败可恢复”。尤其是邮件发送、会议创建、请假写入等副作用动作，一次错误路由可能产生真实业务影响；模型生成的自然语言解释无法替代确定性的权限决策和执行凭证。因此，主 Agent 决策和权限治理不是两个独立插件，而是同一条执行链上的前后约束：前者决定做什么、由谁做，后者决定能不能做、在什么范围内做以及做过之后如何证明。")

    doc.add_heading("1.2 研究目标", level=2)
    add_bullets(doc, [
        "建立可审计的主 Agent 决策入口，把自然语言请求转换为 TaskProfile，并显式表达意图、实体、动作、风险、数据范围、缺失字段和子任务依赖。",
        "建立 AgentCard 与 AgentContract，使候选 Agent 的能力、支持动作、数据范围、风险上限和输入输出契约可被程序校验。",
        "形成“权限硬过滤在前、能力软排序在后”的路由方法，输出包含候选得分、排除原因和 trace_id 的 RoutingDecision。",
        "建立 S-ABAC 四维权限决策、三态审批、Artifact 数据守卫、审计和副作用防重机制，验证 fail-closed 原则。",
        "建设离线评测集与回归测试，量化意图识别、结构化决策和权限阻断效果，并公开当前不足。",
    ])

    doc.add_heading("1.3 专题边界", level=2)
    add_para(doc, "本报告只展开“主 Agent 意图识别与路由”和“权限治理”两部分。多 Agent TaskGraph 编排、Memory、Skill、通用工作流 UI 等模块不作为独立研究对象，仅在说明主 Agent 输出如何进入执行、权限如何覆盖 Agent 与工具时作必要引用。原型使用虚构人员与模拟业务数据，不接入真实银行身份系统或生产业务系统。")
    add_table(doc, ["范围", "纳入内容", "不展开内容"], [
        ["主 Agent", "意图目录、规则/语义识别、TaskProfile、澄清、AgentCard、路由打分、RoutingDecision", "具体 DAG 调度算法、Memory 与 Skill"],
        ["权限治理", "S-ABAC、审批、Agent/Tool PEP、Artifact Guard、审计、Receipt", "真实统一认证、生产密钥管理、跨地域策略中心"],
        ["评测", "离线意图集、合同测试、权限矩阵、治理回归", "线上真实用户 A/B 实验、生产吞吐压测"],
    ], widths=[2.6, 7.2, 6.1])

    doc.add_heading("2 调查研究与方案选型", level=1)
    add_callout(doc, "调研方法", "从语义覆盖、确定性、可解释性、安全边界、工程成本、离线可复现和生产迁移七个维度比较方案。选型不追求单一技术“包办所有判断”，而是明确模型与确定性组件的职责边界。", PALE_BLUE, BLUE)

    doc.add_heading("2.1 主 Agent 意图识别方案比较", level=2)
    doc.add_heading("2.1.1 方案 A：纯关键词/规则", level=3)
    add_para(doc, "该方案通过关键词、正则、词典和有限状态规则识别意图与实体。优势是低成本、低时延、输出稳定、可离线运行，适合“工资、请假、天气、发送”等边界清晰的高频办公意图；规则还可以精确处理否定词、显式动作和确定性风险升级。缺点是同义改写、长句、隐含前置动作和跨句上下文会导致规则数量快速膨胀，维护过程容易产生关键词抢占与顺序依赖。")
    add_para(doc, "本项目的离线规则模式验证了这一特征：主意图准确率、子意图召回率和实体字段准确率较高，但严格整例通过率只有 60.61%，失败主要出现在同义表达、知识查询与报告生成的关键词边界，以及模糊请求的置信度和澄清判断。因此纯规则可以作为确定性基线和降级路径，但不宜作为复杂自然语言的唯一方案。")

    doc.add_heading("2.1.2 方案 B：纯 LLM 分类或直接生成计划", level=3)
    add_para(doc, "纯 LLM 方案把用户请求、候选 Agent 描述和输出格式一次性放入 Prompt，由模型直接给出意图、Agent 或完整计划。其语义泛化能力强，能理解同义表达和隐含目标，开发初期也较快。但模型输出受提示词、模型版本、温度、上下文长度和候选描述变化影响，可能出现格式漂移、虚构 Agent、遗漏否定条件或把“咨询如何发送”误判为“立即发送”。")
    add_para(doc, "项目基线曾出现 Publisher 未遵守约定 JSON 协议、直接回答用户的问题，说明 Prompt 只能影响模型行为，不能构成安全边界。更重要的是，如果让同一个模型同时负责意图理解、权限判断和执行选择，提示注入、模型幻觉或过度代理都可能扩大权限，OWASP 也把 Prompt Injection 与 Excessive Agency 列为 LLM 应用的重要风险[10]。因此本项目不采用“LLM 说允许就允许”的设计，LLM 输出必须进入结构化 Schema，并且只能在权限引擎提供的合法候选集合中发挥作用。")

    doc.add_heading("2.1.3 方案 C：向量召回", level=3)
    add_para(doc, "向量方案把用户任务与 Agent 描述、意图说明或历史样例编码后计算相似度，适合候选规模较大、描述差异明显且需要 Top-K 召回的系统。其优点是比关键词更能覆盖语义近似表达，比让 LLM 全量阅读 Agent 列表更节省上下文。局限在于相似度只表示语义接近，不等于动作、数据范围和权限满足；阈值也需要样本校准。对于当前约 18 个 Agent 安全画像的原型，单独引入向量数据库会增加依赖和调试成本，收益有限。因此向量召回被保留为候选规模扩大后的前置召回层，而不是当前路由的唯一依据。")

    doc.add_heading("2.1.4 方案 D：规则 + 语义模型融合 + 结构化 TaskProfile", level=3)
    add_para(doc, "最终方案将任务理解拆为意图目录、规则识别、可选语义识别、融合与结构校验四层。规则提供高置信度显式意图、否定和实体基线；语义模型补充同义表达、隐含前置动作和复杂条件；IntentFusion 合并候选并保留 explicit/inferred、evidence、text_span 等来源信息；Pydantic TaskProfile 对字段类型和枚举进行校验。关键字段缺失由 ClarificationAnalyzer 产生具体问题，而不是依赖一个抽象置信度盲目追问。")
    add_table(doc, ["方案", "语义泛化", "确定性", "可解释", "离线可用", "安全边界", "本项目结论"], [
        ["纯规则", "中/低", "高", "高", "高", "可控", "作为基线与降级路径"],
        ["纯 LLM", "高", "低", "中", "低", "不可单独信任", "不直接负责权限与最终执行"],
        ["向量召回", "中/高", "中", "中", "取决于模型", "不能表达权限", "规模扩大后用于 Top-K"],
        ["Planner 直接选 Agent", "高", "低", "低/中", "低", "任务理解与路由耦合", "保留规划，不作为安全路由器"],
        ["融合 + TaskProfile", "高", "高", "高", "规则层可用", "可接硬约束", "当前采用"],
    ], widths=[2.6, 1.6, 1.5, 1.5, 1.7, 2.4, 4.0], font_size=8.7)

    doc.add_heading("2.2 Agent 路由组织方式比较", level=2)
    doc.add_heading("2.2.1 去中心化 handoff", level=3)
    add_para(doc, "去中心化 handoff 允许一个 Agent 把控制权转移给另一个 Agent，适合开放式协作和能力动态发现。AutoGen 强调通过可对话 Agent 组合多种交互模式[5]，LangGraph 也提供 handoff、subagent、router 和 custom workflow 等模式[4]。其问题是跨 Agent 调用链容易变长，权限主体和数据边界难以集中审计；部门 Agent 若能自由发现其他 Agent，还可能绕过主入口扩大权限。对于强调统一入口和合规审计的数字员工，本项目不采用部门 Agent 自由互调。")

    doc.add_heading("2.2.2 中心化主 Agent + 能力卡路由", level=3)
    add_para(doc, "中心化模式要求所有跨部门任务都先经过主 Agent。主 Agent 并不以系统管理员身份替用户执行，而是继承原始用户身份，先获得 authorized_agent_ids，再基于 AgentCard 做能力匹配。LangGraph 的 Router 模式说明了“分类后转入专用流程”的通用性[4]，但本项目在普通 Router 之前增加权限硬过滤，在之后增加 AgentContract 与输入输出 Schema 校验，使路由既可解释又可执行。")
    add_table(doc, ["路由方式", "优点", "主要风险/成本", "适用性判断"], [
        ["去中心化 handoff", "灵活、自治、适合开放探索", "调用链和权限传播难审计", "不适合作为本课题主链路"],
        ["LLM Supervisor", "语义能力强、开发快", "候选漂移、可重复性弱", "只能在合法候选内辅助选择"],
        ["静态规则路由", "稳定、低成本", "新增场景需持续加规则", "适合高确定性入口"],
        ["向量 Top-K + 重排", "可扩展、节省上下文", "需要索引、样本和阈值维护", "Agent 数量扩大后引入"],
        ["AgentCard + 硬过滤 + 打分", "可审计、可校验、权限可前置", "要求维护能力卡和契约", "当前采用"],
    ], widths=[3.2, 4.1, 4.8, 3.6])

    doc.add_heading("2.3 权限控制模型比较", level=2)
    doc.add_heading("2.3.1 ACL 与 RBAC", level=3)
    add_para(doc, "ACL 直接维护“主体—资源—权限”映射，简单直接，但当用户、Agent 和工具数量增长时，配置会碎片化。RBAC 按组织角色聚合权限，NIST 的 RBAC 模型能够降低企业权限管理成本[3]，适合表达 HR 经理、工程师、研究员等稳定岗位。然而数字员工请求还受到数据敏感级别、动作类型、收件人、任务目的、网络环境、时间、金额和是否不可逆等上下文影响，仅靠角色会导致角色爆炸，或把同一角色在不同场景下的权限配置得过宽。")

    doc.add_heading("2.3.2 ABAC、ReBAC 与 S-ABAC", level=3)
    add_para(doc, "NIST SP 800-162 将 ABAC 定义为：根据主体、对象、请求操作以及必要时的环境属性，与策略规则共同决定授权[2]。ABAC 能把部门、岗位、密级、资源敏感级别、动作、时间和环境纳入同一决策，天然适合 Agent/Tool 的细粒度控制。ReBAC 则擅长表达所有者、成员、上下级和共享关系，Google Zanzibar 展示了关系模型在大规模统一授权中的价值[8]，但原型当前没有复杂组织关系图，也不需要全球一致性授权基础设施。")
    add_para(doc, "本项目选择 S-ABAC，将 Scenario 作为明确的一等维度，形成 Subject、Object、Scenario、Action 四元组。Scenario 不只是环境变量，还承载业务目标、任务类型、数据范围、意图标签和 Agent 与任务的适配结果。这使“同一 HR 经理读取员工基本信息”和“同一 HR 经理外发工资证明”能够得到不同决策，同时保留将所有者关系通过 Artifact Guard 单独强制执行的能力。")

    doc.add_heading("2.3.3 XACML、OPA 与 Cedar", level=3)
    add_para(doc, "XACML 提供标准化的属性访问控制语言与 PDP/PEP 体系，表达能力完整，但 XML 策略较重[7]。OPA 使用 Rego 将策略决策从业务代码中解耦，支持以结构化 JSON 作为输入[6]，适合微服务、网关和 Kubernetes 等通用场景；Cedar 以 principal、action、resource、context 四部分描述授权请求，并强调 Schema 验证和显式 forbid 优先[9]。三者都为生产化外置策略中心提供了参考。")
    add_para(doc, "原型阶段没有直接引入 OPA/Cedar 服务，原因不是其能力不足，而是当前目标是验证数字员工场景属性、审批状态和执行闸门，嵌入式 Python PolicyEngine 更便于离线演示、断点调试和与现有 FastAPI/调度器集成。同时实现按照 PIP、PDP、PEP、PAP 分层，并把授权输入结构化，后续可把当前 PDP 替换成 OPA、Cedar 或企业权限平台，而无需重写 AgentProxy 和 Tool Wrapper 的执行点。")
    add_table(doc, ["模型/引擎", "核心依据", "优势", "局限", "本项目定位"], [
        ["ACL", "用户—资源列表", "直观简单", "规模大时难维护", "不单独采用"],
        ["RBAC", "用户角色", "符合组织结构", "难表达动态场景", "作为 Subject 属性"],
        ["ABAC", "主体/对象/动作/环境属性", "细粒度、动态", "属性和策略治理成本高", "基础模型"],
        ["ReBAC", "主体与资源关系", "适合所有权、共享与层级", "需要关系图和一致性设施", "Artifact 所有权局部采用"],
        ["XACML", "标准化属性策略", "规范完整、互操作", "XML 较重", "体系结构参考"],
        ["OPA/Rego", "通用策略即代码", "业务与策略解耦", "增加运行服务与 Rego 成本", "生产迁移候选"],
        ["Cedar", "Principal/Action/Resource/Context", "Schema 与 forbid 语义清晰", "需引入新策略栈", "生产迁移候选"],
        ["S-ABAC", "Subject/Object/Scenario/Action", "贴合 Agent 业务场景、可解释", "当前为原型内嵌实现", "当前采用"],
    ], widths=[2.3, 3.1, 3.5, 3.7, 3.4], font_size=8.6)

    doc.add_heading("2.4 审批、数据传递与副作用方案比较", level=2)
    add_para(doc, "权限决策如果只有允许和拒绝两态，会出现两种极端：为了便利而放宽高风险动作，或为了安全把所有敏感业务完全阻断。本项目增加 REVIEW_REQUIRED，表示主体具备静态资格，但本次具体动作必须由人确认。审批签名绑定稳定场景事实和调用参数，批准后只允许消费一次；DENY 则不能通过人工审批被升级为允许。")
    add_para(doc, "步骤间直接传递原始字典虽然方便，但难以表达所有者、敏感级别、来源和完整性。项目改为 Artifact 与 ArtifactRef：payload 存在受保护目录，checkpoint 只保存引用、校验和和敏感级别。Artifact Guard 先做所有权判断，再调用 PolicyEngine；跨用户访问必须在可信来源中显式授权。对于邮件发送等副作用，普通自动重试可能造成重复操作，因此通过 task_id、step_id 和标准化输入生成幂等键，原子写入 STARTED Receipt，成功后记录外部操作号。恢复时已成功直接复用，状态不确定则进入人工对账。")
    add_table(doc, ["治理问题", "简单方案", "简单方案风险", "选用方案"], [
        ["高风险但具备资格", "一律允许/一律拒绝", "过宽或业务不可用", "三态决策 + 一次性审批"],
        ["跨步骤敏感数据", "共享 State 原文", "所有权、血缘和泄露难控制", "ArtifactRef + Guard"],
        ["邮件/会议重试", "失败自动重试", "超时后可能重复发送/创建", "幂等键 + 持久化 Receipt"],
        ["异常或配置缺失", "降级为默认允许", "未知状态下扩大权限", "Fail Closed"],
    ], widths=[3.0, 3.2, 4.7, 5.0])

    doc.add_heading("2.5 最终选型结论", level=2)
    add_callout(doc, "最终技术路线", "主 Agent：Intent Catalog + 规则识别 + 可选语义识别 + TaskProfile + Clarification + AgentCard + 权限硬过滤 + 确定性打分。权限治理：S-ABAC + 四级 PEP + ALLOW/DENY/REVIEW_REQUIRED + 一次性审批 + Artifact Guard + Receipt + 审计。", PALE_GOLD, "C9782B")

    doc.add_heading("3 总体技术架构", level=1)
    add_figure(doc, figures["architecture"], "图 3-1  主 Agent 决策与权限治理总体架构", width=16.0)
    add_para(doc, "系统统一入口为 make_routing_decision。请求首先结合用户身份、会话和上下文生成 TaskProfile；运行时 Agent 注册信息被转为 AgentCard；权限层提前计算 authorized_agent_ids，使路由器看不到无权候选；确定性路由器按意图、能力、场景、数据范围、历史基线和可用性打分，输出 RoutingDecision。真正派发 Agent、调用 Tool 和读取 Artifact 时分别再次鉴权，防止计划生成后上下文变化或内部工具绕过。")
    add_para(doc, "架构采用“模型与策略解耦”：模型输出只作为事实候选，PolicyEngine 和 Schema 校验才决定能否进入执行。这样既保留语言模型对复杂表达的理解能力，又使安全结果不依赖模型是否听从 Prompt。所有关键对象携带 task_id、decision_id、trace_id、approval_id 或 external_op_id，形成从用户请求到外部副作用的证据链。")
    add_table(doc, ["层级", "主要组件", "输入", "输出/责任"], [
        ["任务理解层", "IntentRecognizer、TaskProfiler、ClarificationAnalyzer", "自然语言、会话上下文", "TaskProfile、澄清问题"],
        ["候选与路由层", "AgentCard、DepartmentRouter", "TaskProfile、Agent 注册信息", "候选得分、排除原因、RoutingDecision"],
        ["权限决策层", "PIP、PolicyEngine/PDP", "Subject/Object/Scenario/Action", "ALLOW / DENY / REVIEW_REQUIRED"],
        ["执行控制层", "Agent PEP、Tool Wrapper、ApprovalStore", "路由结果、调用参数", "受控派发、审批暂停/恢复"],
        ["数据与副作用层", "Artifact Guard、ReceiptStore、Audit", "ArtifactRef、幂等键", "受控读取、防重复、审计证据"],
    ], widths=[2.5, 4.1, 4.0, 5.3])

    doc.add_heading("4 主 Agent 意图识别与权限感知路由设计", level=1)
    doc.add_heading("4.1 TaskProfile：把自然语言变成稳定业务对象", level=2)
    add_para(doc, "TaskProfile 是主 Agent 的核心中间表示。它把“模型觉得用户想做什么”转换成程序可验证的字段，包括主意图、子意图、主目标、业务动作、操作模式、实体、数据范围、预期能力、风险等级、不可逆标志、缺失字段、置信度、意图节点、子任务和依赖。意图与主目标被区分：例如“帮李娜开收入证明并发给王经理”的第一个执行意图可能是员工信息查询，但用户的主目标是生成并发送证明。")
    add_code(doc, '''{
  "task_id": "task_income_proof_001",
  "intent": "employee_information_query",
  "primary_goal_intent": "document_generation",
  "sub_intents": [
    "employee_information_query", "salary_query",
    "document_generation", "message_or_email_send"
  ],
  "entities": {
    "employee_name": "李娜", "recipient": "王经理",
    "document_type": "income_proof"
  },
  "action": "send",
  "risk_level": "HIGH",
  "irreversible": true,
  "missing_fields": []
}''', "示例 4-1  收入证明发送任务的 TaskProfile 摘要")

    doc.add_heading("4.2 分层意图识别流程", level=2)
    add_figure(doc, figures["decision"], "图 4-1  意图识别、结构校验与决策门流程", width=14.2)
    add_numbered(doc, [
        "意图目录统一定义受支持的意图、关键词、动作、风险和能力映射，避免多个模块各自维护字符串。",
        "规则识别器对明确关键词、否定词、人物、时间、文档类型和发送对象进行稳定识别，并生成带 evidence 与 text_span 的候选。",
        "语义识别器在配置模型时处理同义改写、隐含前置动作、条件任务和模糊指代；未配置或调用失败时可退化到规则结果。",
        "IntentFusion 合并规则与语义候选，去重并保留 provenance，否定意图只用于审计，不能进入可执行 subtasks。",
        "TaskProfiler 补充业务依赖、动作、数据范围和风险；ClarificationAnalyzer 根据字段契约判断是否必须追问。",
        "Pydantic Schema 校验结构，不能解析、缺少关键对象或存在歧义时输出 CLARIFY，而不是猜测。",
    ])

    doc.add_heading("4.3 澄清与置信度的职责分离", level=2)
    add_para(doc, "项目不把“路由得分不够高”直接等同于“用户信息缺失”。是否澄清由业务字段契约决定，例如发送动作缺少 recipient、会议缺少明确参会对象时必须追问；Agent 匹配得分只是候选能力信号。当前实现中，存在 missing_fields 或 needs_clarification 时输出 CLARIFY；否则最高候选分数不低于 0.80 时为高置信派发，0.55～0.80 为有能力候选派发，低于 0.55 则拒绝。复合任务如果 Top-K 对子任务能力覆盖率达到 0.80，也可以进入派发。")
    add_callout(doc, "为什么要区分", "用户信息不完整时应提出具体问题；Agent 得分稍低时如果仍存在明确、有契约且授权的候选，继续追问用户并不能改善 Agent 能力。把两者分开可以减少无意义澄清。", PALE_TEAL, TEAL)

    doc.add_heading("4.4 AgentCard 与 AgentContract", level=2)
    add_para(doc, "AgentCard 将 Agent 名称之外的能力信息标准化，包括部门、capabilities、intents、supported_actions、accepted_data_scopes、scenario_tags、risk_ceiling、required_grants、tool_scopes、状态和版本。AgentContract 进一步声明 requires、produces、Schema 引用和基数，使“语义上像是合适”升级为“输入输出合同可验证”。只有来自可信 Registry 且存在规划契约的 Agent 才应进入正式规划候选。")
    add_table(doc, ["字段", "作用", "路由/治理意义"], [
        ["capabilities / intents", "声明业务能力与支持意图", "计算能力、意图和复合覆盖率"],
        ["supported_actions", "read/write/generate/send 等", "阻断动作不兼容候选"],
        ["accepted_data_scopes", "可处理的数据范围", "减少跨域数据暴露"],
        ["risk_ceiling", "Agent 可承担的最高风险", "高风险任务不能派给低风险 Agent"],
        ["required_grants", "访问所需授权", "形成 RoutingDecision 的授权说明"],
        ["requires / produces", "输入输出 Artifact 合同", "计划与运行时 Schema 校验"],
        ["version / status", "版本和在线状态", "可用性过滤与审计复现"],
    ], widths=[3.7, 5.0, 6.9])

    doc.add_heading("4.5 权限硬过滤与确定性打分", level=2)
    add_para(doc, "路由按两阶段进行。第一阶段是硬过滤：用户无权访问、Agent 离线、动作不支持、风险超过上限的候选直接进入 excluded_agents，并记录 PERMISSION_DENIED、AGENT_UNAVAILABLE、ACTION_UNSUPPORTED 或 RISK_CEILING_EXCEEDED。第二阶段只对合法候选计算软分数。权限从不进入加权公式，因此一个无权 Agent 即使语义分数为 1.0，也没有进入候选排序的机会。")
    add_code(doc, '''RouteScore =
  0.35 × intent_match
+ 0.25 × capability_match
+ 0.15 × scenario_match
+ 0.10 × data_scope_match
+ 0.10 × history_score
+ 0.05 × availability_score''', "公式 4-1  当前确定性路由评分")
    add_para(doc, "当前 history_score 使用 0.5 基线、availability_score 对在线 Agent 取 1.0，说明历史成功率和成本时延尚未接入真实运行数据；报告不把这两个占位项表述为已完成的智能优化。对于复合任务，意图、能力和场景得分按各子任务覆盖率计算，避免只匹配主意图而漏掉后续发送或文档能力。")

    doc.add_heading("4.6 RoutingDecision 与可解释性", level=2)
    add_para(doc, "路由结果不是一个 Agent 名称，而是包含 decision_id、task_id、selected_agent、候选列表、score_breakdown、reason_codes、required_grants、excluded_agents 和 trace_id 的结构化对象。页面和审计可以回答“为什么选择它”“还有哪些候选”“哪些候选因权限或能力被排除”“当前决策是派发、澄清还是拒绝”。")
    add_code(doc, '''{
  "decision": "DISPATCH",
  "selected_agent": "RemoteHRAssistantAgent",
  "confidence": 0.94,
  "reason_codes": ["HIGH_CONFIDENCE_ROUTE"],
  "candidate_agents": [{
    "agent_id": "RemoteHRAssistantAgent",
    "score": 0.94,
    "reason_codes": ["AUTHORIZED", "INTENT_MATCH", "CAPABILITY_MATCH"]
  }],
  "excluded_agents": [{
    "agent_id": "RemoteEmailDispatchAgent",
    "reason_code": "PERMISSION_DENIED"
  }]
}''', "示例 4-2  RoutingDecision 摘要")

    doc.add_heading("4.7 主 Agent 决策的安全边界", level=2)
    add_bullets(doc, [
        "TaskProfile 中的业务事实可以来自模型，但身份、授权列表、Agent Contract 和安全属性必须来自受信配置或注册中心。",
        "候选得分只能排序合法候选，不能覆盖 PolicyEngine 的 DENY。",
        "模型生成的 Agent 名称必须存在于 Registry，模型不能临时发明并执行一个 Agent。",
        "咨询、否定和条件意图需保留 provenance；被否定动作不得进入可执行子任务。",
        "任何关键字段缺失或无有能力候选时，不使用通用 Agent 兜底执行敏感动作。",
    ])

    doc.add_heading("5 S-ABAC 权限治理设计", level=1)
    doc.add_heading("5.1 四维属性模型", level=2)
    add_para(doc, "S-ABAC 在标准 ABAC 基础上强化 Scenario。一次授权请求由 Subject、Object、Scenario、Action 共同决定，角色只是 Subject 的一个属性，而不是全部权限依据。当前演示配置包含 6 个用户画像、18 个 Agent 安全属性、42 个资源安全属性和 3 条显式策略；未注册资源和未知用户默认拒绝。")
    add_table(doc, ["维度", "典型属性", "解决的问题"], [
        ["Subject 主体", "user_id、role、department、job_role、clearance、trust、grants", "谁在请求，具有什么资格"],
        ["Object 对象", "Agent/Tool、sensitivity、allowed roles、required grants、operation modes", "请求访问什么资源"],
        ["Scenario 场景", "task_type、business_goal、data_scope、scenario_tags、risk、time/network、fit", "为什么、在什么上下文访问"],
        ["Action 动作", "read/query/write/send/delete、amount、batch_size、irreversible", "准备执行什么以及影响大小"],
    ], widths=[2.8, 7.3, 5.7])
    add_code(doc, '''EffectivePermission =
  UserPermission
∩ AgentBoundary
∩ TaskScope
∩ ToolPolicy
∩ ContextPolicy''', "公式 5-1  有效权限是多重边界的交集")

    doc.add_heading("5.2 PIP、PDP、PEP、PAP 分层", level=2)
    add_para(doc, "PIP 负责提供用户、Agent、工具、任务和环境属性；PDP 由 PolicyEngine 计算策略；PEP 分布在路由、Agent 派发、工具包装器和 Artifact 读取路径；PAP 当前由 Python/JSON 配置承担策略管理。该分层与 XACML/OPA 等标准实践兼容，使原型可以内嵌运行，又保留外置策略引擎的迁移接口。")
    add_table(doc, ["组件", "当前实现职责", "后续生产化方向"], [
        ["PIP", "SecurityContextBuilder、TaskProfile、Registry 安全属性", "接入统一身份、组织、设备与网络属性服务"],
        ["PDP", "PolicyEngine.evaluate", "替换为 OPA/Cedar/企业策略中心"],
        ["PEP", "路由前过滤、enforce_agent_dispatch、SecureToolWrapper、Artifact Guard", "网关与服务端双重执行"],
        ["PAP", "s_abac_config.py 与策略版本", "策略审批、灰度、回滚和冲突检测"],
    ], widths=[2.3, 7.4, 6.2])

    doc.add_heading("5.3 四道治理闸门", level=2)
    add_figure(doc, figures["security"], "图 5-1  路由、派发、工具与结果四道权限闸门", width=15.3)
    add_numbered(doc, [
        "路由前：根据原始用户身份得到 authorized_agent_ids，不允许无权 Agent 进入模型或打分候选。",
        "派发前：enforce_agent_dispatch 重新校验目标 Agent、动作、数据范围、风险和场景，防止计划与执行上下文漂移。",
        "工具调用前：SecureToolWrapper 或远程 Tool Gate 对每次参数化调用重新鉴权；Agent 有权被调用不等于其全部工具都可被调用。",
        "结果返回前：Artifact Resolver 通过所有权、敏感级别、数据范围、校验和和审计控制 payload 读取。",
    ])

    doc.add_heading("5.4 ALLOW、DENY 与 REVIEW_REQUIRED", level=2)
    add_para(doc, "PolicyEngine 默认从 DENY 开始。静态资格与资源约束全部满足且不存在强制复核条件时返回 ALLOW；缺少角色、授权、密级、场景适配或资源未注册时返回 DENY；主体有资格但对象声明 requires_approval、动作不可逆或敏感级别较高时返回 REVIEW_REQUIRED。该三态区分“绝对不允许”和“允许条件尚未完成”。")
    add_table(doc, ["决策", "语义", "工作流行为", "是否可通过审批改变"], [
        ["ALLOW", "满足全部策略", "直接进入受控执行", "无需审批"],
        ["REVIEW_REQUIRED", "静态资格满足，但需人工确认", "保存状态、创建审批、暂停", "批准后仅对绑定操作生效"],
        ["DENY", "主体/资源/场景/动作不满足", "立即终止，不产生副作用", "不能被审批提升"],
    ], widths=[2.7, 5.1, 5.1, 3.1])

    doc.add_heading("5.5 审批签名与一次性消费", level=2)
    add_para(doc, "ApprovalStore 为待审批动作生成稳定签名。签名绑定 task_id、资源、动作、稳定场景事实和调用参数；如果收件人、金额、操作模式或业务参数在批准后被修改，旧审批不能匹配。审批记录通过文件锁进行消费，一个已批准请求在并发恢复时只能被一个执行上下文消费；同一执行上下文可缓存已消费结果，避免一个步骤内部重复检查。被拒绝或过期审批保持不可执行。")
    add_callout(doc, "审批不是万能通行证", "只有 REVIEW_REQUIRED 能进入审批流程；DENY 代表主体本身无资格或违反硬约束，管理员点击“批准”也不能把 DENY 改成 ALLOW。", PALE_RED, "B34E4E")

    doc.add_heading("5.6 Artifact 数据权限与血缘", level=2)
    add_para(doc, "成功步骤的结果被封装为 Artifact，携带 artifact_id、version、owner_user_id、producer_agent_id、sensitivity、checksum 和 derived_from。下游只接收 ArtifactRef，并通过统一 Resolver 读取 payload。所有权守卫先于 PolicyEngine 判断，因为通用策略引擎未直接表达 owner_user_id 与 subject.id 的相等关系；跨用户读取必须在可信的 allowed_reader_ids 中显式授权。敏感级别按任务风险和上游血缘取最大值，避免派生报告把工资数据错误降级为普通数据。")
    add_para(doc, "持久化顺序为“payload 落盘—脱敏 checkpoint 索引—步骤完成”。checkpoint 不保存工资等正文，只保存引用、校验和、logical_name 和 sensitivity。恢复时若文件缺失或校验和不一致，系统 fail closed。审计只记录主体、逻辑名、敏感级别、所有者、允许/拒绝、原因和 cross_user 标志，不记录 payload 或 URI。")

    doc.add_heading("5.7 副作用幂等与 Receipt", level=2)
    add_para(doc, "对邮件发送、会议创建等副作用，幂等键用于标识“这是不是同一次业务操作”，Receipt 用于证明“该操作目前处于什么执行状态”。幂等键由 task_id、step_id 和去除 approval_id、idempotency_key 等易变字段后的标准化业务输入计算 SHA-256。同一任务、同一步骤、同一业务输入在审批恢复或进程重启后得到相同键。")
    add_code(doc, '''idempotency_key = SHA256(
  task_id + "|" + step_id + "|" + normalized_input
)''', "公式 5-2  副作用操作身份")
    add_para(doc, "执行器在真正调用外部系统前通过 claim_if_absent 原子写入 STARTED Receipt 并获得 claim_id。并发执行器看到同一 STARTED 时不得再次执行；成功后只有持有 claim_id 的执行器可以写入 SUCCEEDED、external_op_id 和输出 Artifact。恢复时，可信 SUCCEEDED 直接复用；STARTED 表示外部操作可能已经成功但本地尚未落盘，系统转 NEEDS_RECONCILIATION，而不是盲目重试；Receipt 文件损坏时拒绝执行。")
    add_table(doc, ["Receipt 情况", "系统判断", "处理方式"], [
        ["不存在", "尚无执行声明", "原子写 STARTED，获胜者执行"],
        ["可信 SUCCEEDED", "外部操作已完成", "跳过执行并复用结果"],
        ["STARTED / IN_PROGRESS", "结果不确定或其他实例执行中", "停止自动重试，进入对账"],
        ["损坏或身份字段不一致", "无法证明历史状态", "Fail Closed，不执行"],
        ["确认未产生副作用", "可安全释放", "人工/确定性确认后 release_for_retry"],
    ], widths=[3.4, 5.0, 7.2])

    doc.add_heading("5.8 端到端治理示例", level=2)
    add_para(doc, "以“查询李娜的基本信息和工资，生成收入证明并发给王经理”为例，主 Agent 识别四个意图并形成高风险、不可逆的 TaskProfile。路由前，工程师用户会因为缺少 salary_read 等授权而无法看到工资 Agent；HR 经理可进入人员与工资候选，但工资查询或外发邮件可能返回 REVIEW_REQUIRED；批准后，审批只绑定当前人员、文档和收件人。查询结果作为 restricted Artifact 归属于原始用户，文档 Agent 通过 ArtifactRef 读取，邮件 Agent 只有在工具级再次鉴权后才能发送。发送前写 STARTED Receipt，成功后记录邮件系统 message_id。整个过程能够回答谁发起、为什么路由、为何审批、读取了什么逻辑数据以及外部操作是否完成。")

    doc.add_heading("5.9 原型安全边界", level=2)
    add_bullets(doc, [
        "已实现：S-ABAC、所有权、敏感级别、Schema、校验和、存储隔离、审批、审计、Receipt 和失败关闭。",
        "未实现：真实统一身份认证、生产级密钥管理、AES-GCM 静态加密和 Windows ACL 专用实现。",
        "SHA-256 校验和用于发现意外损坏与简单修改，不能宣称对拥有写权限的攻击者实现密码学防篡改。",
        "当前演示用户、Agent 和资源属性为静态配置；生产环境必须由受信身份、组织和注册中心提供。",
        "原型不接入真实个人信息、工资或生产发送系统，评测数据均为模拟数据。",
    ])

    doc.add_heading("6 评测集、实验方法与量化结果", level=1)
    doc.add_heading("6.1 评测原则", level=2)
    add_para(doc, "评测遵循可离线复现、按能力分层、成功与失败路径并重三个原则。意图识别不只统计主标签，还检查子意图集合、实体、动作、依赖、否定、澄清和置信度区间；权限治理不只验证允许路径，还验证未知资源、岗位不匹配、越权工资、审批前副作用、并发审批消费、跨用户 Artifact 和 Receipt 恢复。")
    add_table(doc, ["评测层", "数据/测试", "数量", "主要指标"], [
        ["意图画像", "tests/intent_eval_cases.json，rule 模式", "33 条", "主意图、子意图、实体、依赖、澄清、否定"],
        ["主 Agent 合同", "意图、澄清、合同规划、TaskGraph 转换、AgentContract", "113 项", "结构与合同断言通过率"],
        ["权限 API 矩阵", "permission_governance_eval.json 自动化用例", "12 条", "ALLOW/DENY/REVIEW_REQUIRED 与辅助字段"],
        ["权限核心治理", "S-ABAC、审批、Artifact、Receipt、恢复、离线 HR Demo", "105 项", "阻断、审批、防重、恢复、审计"],
        ["专题合并回归", "含 Web 治理 API 与静态资源断言", "253 项", "整体回归通过率及失败归因"],
    ], widths=[2.7, 6.2, 2.0, 5.7])

    doc.add_heading("6.2 意图画像评测集设计", level=2)
    add_para(doc, "33 条评测覆盖 regression、challenge 和 semantic 三个层级，包含单意图、跨域复合、隐式依赖、同义词、去重、条件任务、缺失字段、否定、咨询误触发、知识边界和发送对象等类型。每条用例按需标注 expected_primary_intent、expected_sub_intents、expected_entities、expected_subtask_actions、expected_dependency_edges、expected_missing_fields、expected_risk_level、expected_irreversible 和 expected_confidence_range。严格整例只有全部标注同时命中才算通过。")
    add_table(doc, ["代表用例", "输入摘要", "关键验证点"], [
        ["hr_income_proof_send_order", "查信息、生成收入证明、发给经理", "信息→工资→文档→发送依赖顺序"],
        ["semantic_negated_send", "生成报告，但不要发送", "发送可审计但不可执行"],
        ["semantic_consultation_permissions", "了解发送证明需要哪些权限", "咨询不能误触发工资/生成/发送"],
        ["semantic_conditional_meeting", "有空则安排会议并通知", "条件门与依赖不是简单线性顺序"],
        ["missing_recipient_income_proof", "生成证明然后发送", "保留完整前置依赖并追问收件人"],
        ["knowledge_report_boundary", "总结员工请假制度形成材料", "“员工”不能抢占知识查询主意图"],
    ], widths=[4.2, 5.6, 5.8])

    doc.add_heading("6.3 意图画像量化结果", level=2)
    add_figure(doc, figures["metrics"], "图 6-1  33 条离线规则模式意图评测关键指标", width=14.8)
    add_table(doc, ["指标", "结果", "解释"], [
        ["严格整例通过率", "60.61%（20/33）", "所有已标注字段和区间同时满足才通过"],
        ["主意图准确率", "93.94%", "顶层业务域识别较稳定"],
        ["主目标准确率", "100.00%", "已标注主目标用例全部正确"],
        ["子意图 Precision / Recall", "100.00% / 98.48%", "几乎不多执行，少量隐式意图漏召回"],
        ["实体字段准确率", "96.46%", "人员、收件人、文档类型总体稳定"],
        ["依赖完全准确率", "92.31%", "复合任务大部分顺序正确"],
        ["澄清判定准确率", "75.00%", "模糊/否定请求仍需优化"],
        ["否定动作误执行率", "0.00%", "标注的否定动作未进入执行子任务"],
    ], widths=[4.1, 3.4, 8.0])
    add_para(doc, "严格整例通过率低于各关键字段指标并不矛盾：一个用例即使意图、实体和依赖均正确，只要置信度落在标注区间之外或 missing_fields 多出一项，也会整例失败。13 条失败中，多个属于置信度区间和澄清字段偏差；真正的主意图错误主要集中在 knowledge_report_boundary 和 leave_material_paraphrase。该结果说明规则模式已能稳定支撑明确办公任务，但要覆盖弱表达和语义边界，仍需运行并校准 hybrid 模式。")
    add_callout(doc, "诚实报告结果", "本报告没有用 113/113 的单元测试通过率替代真实样例上的 60.61% 严格整例通过率。两者分别回答“实现是否符合已编码合同”和“面对多样自然语言是否完全命中标注”，应同时呈现。", PALE_GOLD, "C9782B")

    doc.add_heading("6.4 权限 API 矩阵结果", level=2)
    add_para(doc, "权限 API 评测使用 FastAPI TestClient 在本地直接执行 12 条自动化用例，避免依赖外部服务。结果 12/12 通过，覆盖 5 条 ALLOW、4 条 DENY 和 3 条 REVIEW_REQUIRED。")
    add_table(doc, ["用例", "主体与资源", "预期/实测", "治理意义"], [
        ["PG-API-001", "admin → 风险 Agent", "ALLOW / ALLOW", "合法 Agent 派发"],
        ["PG-API-002", "hr_manager → 工资工具", "REVIEW / REVIEW", "敏感读取需审批"],
        ["PG-API-003", "engineer → 工资工具", "DENY / DENY", "缺 grant 不可审批绕过"],
        ["PG-API-005", "communication_officer → 邮件工具", "REVIEW / REVIEW", "外发副作用审批"],
        ["PG-API-006", "admin → 未注册工具", "DENY / DENY", "管理员也不能调用未知资源"],
        ["PG-API-007", "hr_manager → 邮件 Agent", "DENY / DENY", "Agent 级职责边界"],
        ["PG-API-009", "guest → researcher", "DENY / DENY", "可见不等于岗位匹配"],
        ["PG-API-012", "communication_officer → 会议工具", "REVIEW / REVIEW", "写操作审批"],
    ], widths=[2.6, 5.1, 3.7, 4.8], font_size=8.8)

    doc.add_heading("6.5 回归测试结果", level=2)
    add_figure(doc, figures["validation"], "图 6-2  主 Agent 与权限治理原型验证结果", width=14.2)
    add_table(doc, ["测试组", "结果", "结论"], [
        ["主 Agent 决策相关", "113/113（100%）", "意图、澄清、合同、AgentCard/TaskGraph 相关断言通过"],
        ["权限治理核心", "105/105（100%）", "S-ABAC、审批、Artifact、Receipt、恢复与离线场景通过"],
        ["权限 API 决策矩阵", "12/12（100%）", "三态决策与辅助字段全部匹配"],
        ["专题合并回归", "252/253（99.60%）", "唯一失败为前端静态资源缓存版本字符串断言"],
    ], widths=[4.1, 3.7, 8.1])
    add_para(doc, "合并回归的唯一失败为 test_demo_static_assets_disable_stale_cache_and_include_resume_fixes：测试期待旧的静态资源版本字符串，而当前 index.html 已更新为新的 security.js 缓存版本。该失败不涉及 PolicyEngine、审批、路由、Artifact 或 Receipt，但说明 UI 资源版本和测试断言需要同步。报告将其保留为已知问题，而不是删除失败用例后宣称全量 100%。")

    doc.add_heading("6.6 误差分析与改进方向", level=2)
    add_table(doc, ["问题", "表现", "原因判断", "改进措施"], [
        ["弱语义/同义表达", "休假材料等表达漏掉员工信息前置意图", "规则词典覆盖有限", "运行 hybrid 对照评测，补充语义样例和融合规则"],
        ["关键词边界", "知识制度+汇报场景主意图顺序偏差", "局部关键词优先级影响主目标", "用 primary_goal 与执行顺序分离，增加边界回归"],
        ["澄清过度/不足", "否定或通用请求出现多余 missing_fields", "字段契约未充分结合否定与咨询", "让 ClarificationAnalyzer 读取意图 provenance 和执行性"],
        ["置信度未校准", "业务字段正确但超出人工标注区间", "当前置信度主要是规则组合分", "基于验证集做可靠性曲线和温度校准"],
        ["历史分占位", "路由 history_score 固定 0.5", "尚未接入运行统计", "引入按 Agent/意图分桶的成功率、时延和成本"],
        ["策略配置静态", "演示用户与资源写在配置中", "原型阶段便于离线复现", "接入 IAM/组织/注册中心，增加策略版本与灰度"],
    ], widths=[3.0, 4.4, 4.2, 4.2], font_size=8.6)

    doc.add_heading("7 原型成果展示（前端截图占位）", level=1)
    add_para(doc, "本章预留前端截图位置。建议在治理模式启动后，以相同演示任务分别使用 admin、hr_manager、engineer 和 communication_officer 身份截图，保证页面中的 TaskProfile、路由原因、权限决策和审批状态能够相互对应。粘贴截图时建议裁掉浏览器无关区域，保持图片宽度约 15～16 cm，并在图注下补充一到两句观察结论。")

    add_screenshot_placeholder(doc, "图 7-1  主 Agent 决策台：TaskProfile 与 RoutingDecision", "建议截图内容：原始请求、主/子意图、实体、风险、候选 Agent 分数、selected_agent、reason_codes。")
    add_para(doc, "观察说明（待补）：说明主 Agent 如何从自然语言中识别业务目标，哪些候选因能力不匹配或权限不足被排除，以及最终派发依据。", indent=False)

    add_screenshot_placeholder(doc, "图 7-2  候选 Agent 与权限排除原因", "建议截图内容：authorized candidates、excluded_agents、PERMISSION_DENIED/ACTION_UNSUPPORTED 等原因码。")
    add_para(doc, "观察说明（待补）：重点展示“权限硬过滤先于软排序”，证明越权 Agent 即使能力匹配也不会进入最终候选。", indent=False)

    add_screenshot_placeholder(doc, "图 7-3  S-ABAC 权限预检查页面", "建议截图内容：当前用户 Subject 属性、目标 Agent/Tool Object 属性、Scenario、Action 与三态决策。")
    add_para(doc, "观察说明（待补）：对比 HR 经理查询工资为 REVIEW_REQUIRED、工程师查询工资为 DENY，并说明两种结果不能互换。", indent=False)

    add_screenshot_placeholder(doc, "图 7-4  审批中心与任务恢复", "建议截图内容：pending/approved/rejected 状态、approval_id、资源、动作、参数摘要和审批后恢复按钮。")
    add_para(doc, "观察说明（待补）：说明审批签名绑定具体操作，批准后一次性消费，审批前不能发生邮件发送等副作用。", indent=False)

    add_screenshot_placeholder(doc, "图 7-5  治理时间线、Artifact 与 Receipt", "建议截图内容：Agent/Tool 权限事件、Artifact 读取允许/拒绝、Receipt 状态、external_op_id 与 trace_id。")
    add_para(doc, "观察说明（待补）：展示从路由、审批到副作用完成的可追溯证据链，并指出审计页面不展示敏感 payload。", indent=False)

    doc.add_heading("8 研究成果、局限与后续工作", level=1)
    doc.add_heading("8.1 已形成的成果", level=2)
    add_bullets(doc, [
        "形成主 Agent 统一决策入口：任务画像 → AgentCard → 权限候选 → RoutingDecision。",
        "形成支持主意图、子意图、实体、依赖、否定、澄清和风险的 TaskProfile 结构。",
        "形成权限感知的确定性路由与可解释原因码，权限硬约束不参与软打分。",
        "形成 Subject–Object–Scenario–Action 四维 S-ABAC 与 Agent/Tool 两级执行拦截。",
        "形成一次性审批、Artifact 所有权与血缘、审计、幂等键和 Receipt 恢复机制。",
        "形成 33 条意图评测、21 条权限治理评测定义，以及对应自动化回归。",
    ])

    doc.add_heading("8.2 当前局限", level=2)
    add_para(doc, "原型已证明技术链路可行，但尚未达到生产系统完备性。语义评测本次采用可离线复现的 rule 模式，未调用外部 LLM 运行 hybrid 对照，因此不能据此宣称混合模式优于规则模式；路由历史成功率和成本仍为占位基线；身份、组织和策略来自静态配置；Artifact 未实现静态加密；权限测试主要基于 Mock 资源和本地文件持久化。")

    doc.add_heading("8.3 后续工作", level=2)
    add_numbered(doc, [
        "在固定模型版本、固定 Prompt 和固定温度下运行 rule、semantic、hybrid 三组对照，统计准确率、时延、调用成本和降级成功率。",
        "扩充到任务书六类办公场景的分层评测集，加入错别字、口语、省略、跨域组合、提示注入和高风险请求。",
        "接入运行统计，按 Agent×Intent 计算成功率、P95 时延和失败类型，替换固定 history_score。",
        "将 S-ABAC 策略迁移到 OPA/Cedar 或企业策略中心，保留现有 PEP，增加策略版本、仿真、冲突检测和灰度发布。",
        "接入真实 IAM/组织目录和短时 DelegationGrant，确保主 Agent 永远以原始用户权限执行。",
        "为 Artifact 引入 AES-GCM 与正式密钥管理，在 Windows/生产环境实现明确的存储 ACL。",
        "建立线上影子评测与人工反馈闭环，但在完成合规评审前不连接真实工资和外发系统。",
    ])

    doc.add_heading("9 结论", level=1)
    add_para(doc, "本研究表明，面向数字员工的主 Agent 不能只是一个“更大的 Prompt”。可靠的任务决策需要把自然语言理解结果转成结构化 TaskProfile，以 AgentCard 和 AgentContract 限定可执行能力，并在权限硬过滤之后进行确定性打分。模型擅长处理表达变化，规则和 Schema 擅长保证结构，权限引擎负责不可绕过的边界，三者应分工而不是相互替代。")
    add_para(doc, "在权限治理方面，S-ABAC 将主体、对象、场景和动作同时纳入决策，比仅按角色控制更适合人员工资、文档生成、会议创建和消息外发等动态办公场景。四道 PEP、三态审批、Artifact 数据守卫和 Receipt 防重把“是否允许”“读到什么”“是否已经执行”分别落到可验证机制。现有评测证明原型在确定性合同和治理路径上已经形成稳定闭环，同时离线自然语言严格整例通过率也揭示了语义边界与阈值校准的真实改进空间。该方案可作为后续数字员工统一主入口、外置策略中心和生产级审计体系的技术基础。")

    doc.add_heading("参考文献", level=1)
    refs = [
        "[1] 《实习生任务——面向数字员工的智能任务执行关键技术研究》，内部课题任务书，2026。",
        "[2] NIST. SP 800-162: Guide to Attribute Based Access Control (ABAC) Definition and Considerations. https://doi.org/10.6028/NIST.SP.800-162",
        "[3] NIST. Role Based Access Control Project and INCITS 359 RBAC Model. https://csrc.nist.gov/Projects/Role-Based-Access-Control",
        "[4] LangChain. LangGraph Workflows and Agents: Routing and Multi-agent Patterns. https://langchain-ai.github.io/langgraph/agents/tools/",
        "[5] Wu Q., et al. AutoGen: Enabling Next-Gen LLM Applications via Multi-Agent Conversation. COLM 2024. https://www.microsoft.com/en-us/research/publication/autogen-enabling-next-gen-llm-applications-via-multi-agent-conversation-framework/",
        "[6] Open Policy Agent. OPA Documentation and Policy Language. https://www.openpolicyagent.org/docs",
        "[7] OASIS. eXtensible Access Control Markup Language (XACML) Version 3.0. https://docs.oasis-open.org/xacml/3.0/xacml-3.0-core-spec-os-en.pdf",
        "[8] Pang R., et al. Zanzibar: Google’s Consistent, Global Authorization System. USENIX ATC 2019. https://www.usenix.org/system/files/atc19-pang.pdf",
        "[9] Cedar Policy Language Reference Guide. https://docs.cedarpolicy.com/",
        "[10] OWASP. Top 10 for Large Language Model Applications: Prompt Injection and Excessive Agency. https://owasp.org/www-project-top-10-for-large-language-model-applications/",
        "[11] SuperAgent 项目源码与文档：src/orchestrator、src/security、src/orchestration、docs/权限治理评测集.md、docs/原型安全边界与演示验收.md。",
    ]
    for ref in refs:
        add_para(doc, ref, indent=False, space_after=4)

    doc.add_heading("附录 A 关键实现索引", level=1)
    add_table(doc, ["能力", "主要文件"], [
        ["主 Agent 统一入口", "src/orchestrator/main_agent.py"],
        ["意图识别与融合", "src/orchestrator/intent_recognition.py、intent_catalog.py"],
        ["任务画像与澄清", "src/orchestrator/task_profiler.py、clarification.py"],
        ["AgentCard 与路由", "src/orchestrator/department_router.py"],
        ["TaskProfile / RoutingDecision", "src/contracts/task_profile.py、routing_decision.py"],
        ["S-ABAC 策略", "src/security/policy.py、context.py、enforcement.py"],
        ["Agent/Tool 执行点", "src/security/tool_wrapper.py、remote_tool_gate.py"],
        ["审批", "src/security/approval.py"],
        ["Artifact 权限", "src/orchestration/artifact_guard.py、artifact_payload_store.py"],
        ["幂等与 Receipt", "src/orchestration/completion.py、scheduler.py"],
        ["评测", "tests/intent_eval_cases.json、tests/evaluations/permission_governance_eval.json"],
    ], widths=[5.0, 10.8])

    doc.add_heading("附录 B 评测复现命令", level=1)
    add_code(doc, '''# 1. 离线规则模式意图评测
.\\.venv\\Scripts\\python.exe tests\\evaluate_intent_profile.py \\
  --mode rule --output output\\research_report_intent_eval.md

# 2. 主 Agent 决策相关测试
.\\.venv\\Scripts\\python.exe -m pytest -q -o "addopts=" \\
  tests\\test_intent_recognition.py tests\\test_clarification_analyzer.py \\
  tests\\test_contract_planning.py tests\\test_plan_to_task_graph.py \\
  tests\\test_agent_contract.py

# 3. 权限治理核心测试
.\\.venv\\Scripts\\python.exe -m pytest -q -o "addopts=" \\
  tests\\test_s_abac.py tests\\test_governance_approval.py \\
  tests\\test_artifact_guard.py tests\\test_receipt_persistence.py \\
  tests\\test_scheduler_recovery.py tests\\test_demo_hr_salary_offline.py''')

    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    result = build_report()
    print(result)
