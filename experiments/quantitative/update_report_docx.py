from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[2]
REPORT = ROOT / "doc-yz" / "多智能体协同与执行编排机制调研及原型成果报告-修订版.docx"
OUTPUT = ROOT / "doc-yz" / "多智能体协同与执行编排机制调研及原型成果报告-修订版-量化实验更新.docx"


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def format_table(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            tr_pr = row._tr.get_or_add_trPr()
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)
        for column_index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[column_index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                if column_index > 0 and len(paragraph.text) < 28:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    if row_index == 0:
                        run.bold = True


def keep_table_rows_together(table) -> None:
    for row in table.rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


def add_paragraph_before_anchor(document, anchor, text: str, style: str):
    paragraph = document.add_paragraph(text, style=style)
    paragraph._p.getparent().remove(paragraph._p)
    anchor._p.addprevious(paragraph._p)
    return paragraph


def add_table_before_anchor(document, anchor, headers, rows, widths):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table"
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    format_table(table, widths)
    table._tbl.getparent().remove(table._tbl)
    anchor._p.addprevious(table._tbl)
    return table


def replace_mcp_table(document) -> None:
    table = document.tables[4]
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    values = [
        ("评测工具数 / 唯一案例数", "48 / 53"),
        ("平均过滤前 / 过滤后候选数", "48.00 / 2.17"),
        ("候选压缩率", "95.48%"),
        ("Tool Top-1 / Top-3", "100% / 100%"),
        ("参数 Schema 有效率", "100%"),
        ("应弃选案例弃选率", "100%"),
        ("不存在工具幻觉率", "0%"),
        ("audit 建议与实际契约一致率", "100%"),
        ("选择时延（均值 / P95）", "1.483 ms / 2.025 ms"),
    ]
    for label, result in values:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = result
    format_table(table, [4100, 2700])


def main() -> None:
    document = Document(REPORT)
    replace_mcp_table(document)
    anchor_32 = next(p for p in document.paragraphs if p.text.strip().startswith("3.2 "))
    add_paragraph_before_anchor(document, anchor_32, "3.1.1 结构化计划量化验证", "Heading 3")
    planning_table = add_table_before_anchor(
        document,
        anchor_32,
        ["指标", "结果"],
        [
            ("评测规模", "60 个唯一计划：12 个合法计划、8 类共 48 个非法计划"),
            ("合法计划通过", "12/12，100%"),
            ("非法计划拦截", "48/48，100%"),
            ("误拒 / 漏放", "0 / 0"),
            ("流水线时延（均值 / P95）", "0.902 ms / 1.957 ms；30 轮共 1800 个样本"),
        ],
        [3900, 4600],
    )
    keep_table_rows_together(planning_table)
    add_paragraph_before_anchor(
        document,
        anchor_32,
        "非法计划覆盖任务缺失、重复覆盖、依赖不一致、意图不一致、未知依赖、环依赖、输出绑定无效和"
        "计划快照漂移 8 类错误。结果表明，在固定结构化输入上，TaskProfile、TaskGraph 转换和 "
        "PlanSnapshot 能够形成失效关闭的执行门禁。该实验采用确定性校验与错误注入，衡量的是候选计划"
        "进入执行层之前的结构化治理能力，不等同于大语言模型生成候选计划本身的准确率。",
        "Body Text",
    )

    anchor_33 = next(p for p in document.paragraphs if p.text.strip().startswith("3.3 "))
    add_paragraph_before_anchor(document, anchor_33, "3.2.1 真实 HTTP 协作验证", "Heading 3")
    agent_table = add_table_before_anchor(
        document,
        anchor_33,
        ["检查项", "实验记录"],
        [
            ("执行方式", "三 Agent 年假任务，真实 HTTP 远程调用"),
            ("执行轮次", "连续 5 轮，5/5 完成"),
            ("依赖与并行", "HR Agent 与 Knowledge Agent 并行；Report Agent 等待两路结果汇合"),
            ("Artifact 输出", "每轮均生成 employee.info、policy.info 和 report.markdown"),
            ("证据检查", "Schema、Artifact 血缘和工具调用日志均满足验收条件"),
        ],
        [3000, 5500],
    )
    keep_table_rows_together(agent_table)
    add_paragraph_before_anchor(
        document,
        anchor_33,
        "该结果验证了当前年假案例中基于 depends_on 的并行调度、基于 ArtifactRef 的正式结果传递和"
        "多路 Artifact 汇合。其样本量为 5 个真实 HTTP 执行实例，不能外推为任意 Agent 数量和开放业务"
        "环境下的生产成功率。包含 Approval 与 Email 的五 Agent 扩展流程见 3.5.1。",
        "Body Text",
    )

    mcp_anchor = next(p for p in document.paragraphs if p.text.strip().startswith("当前固定评估集包含"))
    add_paragraph_before_anchor(document, mcp_anchor, "3.3.1 固定评估集 audit 结果", "Heading 3")
    mcp_anchor.text = (
        "当前固定评估集包含 48 个通过 MCP tools/list 获取的工具，以及 53 个唯一案例：48 个普通覆盖"
        "案例、1 个歧义案例、2 个输入不足案例和 2 个未知任务案例。其中 49 个案例具有期望工具标签，"
        "4 个案例期望弃选。30 轮共形成 1590 个进程内时延样本。"
    )
    mcp_anchor.style = "Body Text"
    mcp_result = next(p for p in document.paragraphs if p.text.strip() == "上述数据用于验证当前固定测试集上的机制可行性，不代表开放业务环境中的泛化准确率。")
    mcp_result.text = (
        "固定评估集上，Server 识别、Tool Top-1、Tool Top-3、参数 Schema 有效率、应弃选案例弃选率"
        "以及 audit 建议与实际契约一致率均为 100%，不存在工具幻觉率为 0%。平均候选数从 48.00 降至"
        "2.17，候选压缩率为 95.48%；平均选择时延为 1.483 ms，P95 为 2.025 ms。"
    )
    mcp_result.style = "Body Text"
    mcp_boundary = next(p for p in document.paragraphs if p.text.strip().startswith("当前 Tool Resolver 主要以 audit 模式"))
    mcp_boundary.text = (
        "上述结果验证的是统一发现、分层过滤、评分、Top-K 和 ABSTAIN 在当前固定数据集上的控制面能力。"
        "Tool Resolver 仍主要以 audit 模式接入执行链路，用于给出工具推荐并与 Executor 实际选择结果"
        "进行对照，尚未强制替换 Executor 实际可见工具集合，因此不能据此宣称已经降低了生产环境中的"
        "真实工具误选率。"
    )
    mcp_boundary.style = "Body Text"

    anchor_35 = next(p for p in document.paragraphs if p.text.strip().startswith("3.5 "))
    add_paragraph_before_anchor(document, anchor_35, "3.4.1 有界恢复策略量化对比", "Heading 3")
    recovery_table = add_table_before_anchor(
        document,
        anchor_35,
        ["策略", "全场景自动闭环率", "全场景首次失败恢复率", "平均逻辑调用数", "P95 虚拟时延"],
        [
            ("B0：无自动恢复", "12.4%", "0%", "1.000", "120 ms"),
            ("B1：同 Agent 有界重试", "27.2%", "16.9%", "1.476", "220 ms"),
            ("B2：重试后等价 Agent 改派", "40.0%", "31.5%", "1.604", "360 ms"),
        ],
        [2500, 1250, 1800, 1500, 1450],
    )
    keep_table_rows_together(recovery_table)
    add_paragraph_before_anchor(
        document,
        anchor_35,
        "注：全场景自动闭环率以每种策略的全部 500 次试验为分母，其中包含不可重试、缺少可信备份及"
        "必须进入人工对账的场景，不等同于系统的一般任务成功率。全场景首次失败恢复率以发生首次失败的"
        "试验为分母。副作用结果不确定时进入 NEEDS_RECONCILIATION 属于预期的安全处置结果。",
        "Body Text",
    )
    add_paragraph_before_anchor(
        document,
        anchor_35,
        "实验覆盖 5 类故障、3 种恢复策略，共 1500 次试验。瞬时超时下，B1 和 B2 的完成率及首次"
        "失败后恢复率均达到 100%；主 Agent 持续失败时，只有允许可信等价 Agent 改派的 B2 达到 "
        "100% 完成。不可重试错误、缺少可信备份和副作用结果不确定三类场景均未被错误闭环；其中副作用"
        "结果不确定的 300 条试验全部进入 NEEDS_RECONCILIATION。全部试验中重复副作用和治理违规均为 0。",
        "Body Text",
    )
    add_paragraph_before_anchor(
        document,
        anchor_35,
        "结果同时显示，恢复能力提升伴随着逻辑调用数和虚拟时延增加。该实验采用 Stub Router/Executor、"
        "固定故障模型和虚拟成本，验证的是恢复机制在受控条件下的行为，不代表生产环境恢复率或真实墙钟时延。",
        "Body Text",
    )

    anchor_4 = next(p for p in document.paragraphs if p.text.strip() == "4 总结")
    add_paragraph_before_anchor(document, anchor_4, "3.5.1 五 Agent 扩展流程负面结果", "Heading 3")
    add_paragraph_before_anchor(
        document,
        anchor_4,
        "包含 Approval 和 Email 的五 Agent 流程暴露出审批身份边界：当前验收脚本使用 admin 执行身份，"
        "而现有 S-ABAC 策略允许管理员绕过强制复核，导致 Email 步骤直接进入 SUCCEEDED，未按脚本预期"
        "进入 APPROVAL_REQUIRED。批准和拒绝两个批次原计划各运行 3 轮，但均采用失败即停，实际各执行 "
        "1 轮即停止。因此，该记录用于说明审批策略与验收身份不一致，不能计算或表述为“0/3 成功率”。",
        "Body Text",
    )
    add_paragraph_before_anchor(document, anchor_4, "3.5.2 实验复现说明", "Heading 3")
    add_paragraph_before_anchor(
        document,
        anchor_4,
        "上述实验以 Git 提交 2c9c684 为代码基线，于 2026 年 8 月 11 日执行，产物按运行标识 "
        "20260811-2c9c684 保存，包括案例明细 CSV、汇总 JSON/Markdown、测试报告和图表。最终相关"
        "回归测试为 44 项通过，用时 5.24 s。各比例均使用相应实验样本作为分母；时延分别注明进程内"
        "开销或虚拟故障成本。本章结论应限定为当前代码基线和评测设置下的原型验证结果。",
        "Body Text",
    )

    document.save(OUTPUT)


if __name__ == "__main__":
    main()
