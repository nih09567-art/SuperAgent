from __future__ import annotations

from pathlib import Path

from docx import Document

from update_report_docx import format_table, keep_table_rows_together


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "doc-yz" / "多智能体协同与执行编排机制调研及原型成果报告-修订版-量化实验更新.docx"
OUTPUT = ROOT / "doc-yz" / "多智能体协同与执行编排机制调研及原型成果报告-修订版-MCP选择器优化.docx"


def replace_rows(table, rows: list[tuple[str, ...]], widths: list[int]) -> None:
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    format_table(table, widths)
    keep_table_rows_together(table)


def replace_paragraph(document, starts_with: str, text: str) -> None:
    paragraph = next(p for p in document.paragraphs if p.text.strip().startswith(starts_with))
    paragraph.text = text
    paragraph.style = "Body Text"


def main() -> None:
    document = Document(SOURCE)

    replace_rows(
        document.tables[4],
        [
            ("统一评测中的 Planner→TaskGraph", "5/5 符合金标准"),
            ("统一评测中的非法候选计划拒绝", "1/2；ORCH-007 未拦截缺少 fan-in 来源的计划"),
            ("确定性门禁回归", "60 个唯一计划：12/12 合法计划接受，48/48 非法计划拒绝"),
            ("门禁进程内时延（均值 / P95）", "1.445 ms / 3.795 ms；30 轮共 1800 次调用"),
        ],
        [3900, 4600],
    )
    replace_paragraph(
        document,
        "非法计划覆盖任务缺失",
        "统一评测中，5 个由 Planner 生成并进入 TaskGraph 的案例均符合金标准；但 2 个专门构造的非法候选计划只正确拒绝 1 个。"
        "ORCH-007 删除了一个报告步骤所需的 fan-in 来源，转换器仍予以接受，说明当前门禁能够检查图形结构与绑定形式，却未完整校验金标准要求的来源集合。"
        "另以 60 个固定结构化计划进行确定性回归，12 个合法计划全部接受、8 类共 48 个非法计划全部拒绝。后者验证的是既定错误模式，不等同于大语言模型规划准确率。",
    )

    replace_rows(
        document.tables[5],
        [
            ("统一评测规模", "8 个实际多 Agent 执行案例，覆盖年假、天气差旅、企业风险和收入证明"),
            ("金标准行为符合", "7/8，87.5%"),
            ("真实 HTTP 证据", "年假三 Agent 与五 Agent 场景逐案执行，核对终态、依赖、Artifact 和审批门禁"),
            ("确定性复合 Harness", "补充跨业务 DAG、Artifact、血缘、故障传播和敏感数据断言"),
            ("不符合项", "ORCH-017：薪资明文进入 Scheduler 事件数据"),
        ],
        [3000, 5500],
    )
    replace_paragraph(
        document,
        "该结果验证了当前年假案例中",
        "统一评测不再以同一个年假任务连续运行 5 次作为总体成功率，而是使用 8 个实际多 Agent 执行案例，覆盖 4 类业务工作流。"
        "其中 7 个符合预先冻结的终态、依赖、Artifact、Schema、血缘和安全断言；ORCH-017 因薪资明文进入 Scheduler 事件数据而不符合。"
        "部分案例采用真实 HTTP 远程 Agent，部分跨业务案例采用确定性复合 Harness，因此该指标应称为金标准行为符合率，不能外推为开放生产环境成功率。",
    )

    replace_rows(
        document.tables[6],
        [
            ("冻结盲测集", "36 例：24 个自然语言正例、8 个易混淆例、4 个应弃选例"),
            ("Tool Top-1 / Top-3", "93.75% / 100%"),
            ("工具族宏平均 Top-1", "95.238%"),
            ("应弃选 / 错误弃选", "4/4 / 0/32"),
            ("Schema 缺参阻断", "8/8，100%"),
            ("平均候选数 / 压缩率", "48.000 → 4.917 / 89.757%"),
            ("隔离 Top-K 约束验证", "8/8 成功约束；期望工具实际调用 7/8"),
            ("MCP 聚焦回归", "31 项通过"),
        ],
        [4100, 2700],
    )
    replace_paragraph(
        document,
        "当前固定评估集包含 48 个通过",
        "本次先冻结 36 条盲测案例及其期望结果，再读取 Selector 实现并执行评测；案例包含自然语言正例、易混淆例和应弃选例。"
        "工具注册表固定为 48 个 MCP Tool，并记录了评测集与 Registry 的 SHA-256，避免测试后修改金标准。",
    )
    replace_paragraph(
        document,
        "固定评估集上，Server 识别",
        "首轮盲测发现 32 个应选择工具的案例中有 6 个被错误弃选。针对否定语义未识别、推断结果硬过滤、场景关键词冲突，"
        "以及被 Schema 淘汰的高分工具压制合法候选等原因优化后，Tool Top-1 由 75.0% 提升至 93.75%，Top-3 由 84.375% 提升至 100%，"
        "工具族宏平均 Top-1 由 73.81% 提升至 95.238%。4 个应弃选案例仍全部弃选，错误弃选由 6/32 降至 0/32。"
        "候选数由 48.000 降至 4.917，压缩率为 89.757%；8 个缺参探针仍全部被 Schema 阻断。",
    )
    replace_paragraph(
        document,
        "上述结果验证的是统一发现",
        "生产执行链路暂未启用强制工具集合约束，Executor 仍按原工具集合执行；隔离实验仅验证了 Top-K 强制约束方案的可行性。"
        "隔离的 RecordingExecutor 只读实验中，8/8 成功将工具集约束到 Top-K，7/8 调用了期望工具。"
        "因此，当前结果用于说明分层过滤与排序的原型能力，不能据此宣称已经降低生产环境真实误选率。",
    )

    replace_rows(
        document.tables[7],
        [
            ("B0：无自动恢复", "12.4%", "0%", "1.000", "120 ms"),
            ("B1：同 Agent 有界重试", "27.2%", "16.9%", "1.476", "220 ms"),
            ("B2：重试后等价 Agent 改派", "40.0%", "31.5%", "1.604", "360 ms"),
            ("统一评测：受控故障行为", "7/8", "—", "—", "真实墙钟仅留档"),
        ],
        [2500, 1250, 1800, 1500, 1450],
    )
    replace_paragraph(
        document,
        "实验覆盖 5 类故障、3 种恢复策略",
        "统一评测中的 8 个受控故障案例有 7 个符合金标准；ORCH-016 将首个未执行的下游消费者标记为 FAILED，而金标准要求 SKIPPED。"
        "此外保留 5 类故障、3 种策略、共 1500 次离线对照：瞬时超时下 B1 和 B2 均可恢复，持续失败时仅 B2 的可信改派能够闭环；"
        "副作用结果不确定的 300 条试验全部进入 NEEDS_RECONCILIATION，重复副作用和治理违规均为 0。",
    )
    replace_paragraph(
        document,
        "结果同时显示，恢复能力提升",
        "B0/B1/B2 的全场景自动闭环率分别为 12.4%、27.2% 和 40.0%，表示固定故障模型下的策略敏感性，而不是系统一般任务成功率。"
        "该离线实验使用 Stub Router/Executor、固定故障模型和虚拟成本；真实场景小样本只用于检查预期安全行为，不据此作性能优劣结论。",
    )

    replace_paragraph(
        document,
        "包含 Approval 和 Email 的五 Agent 流程暴露出",
        "新版统一评测不再把早期五 Agent 批次的失败即停记录换算为“0/3”。在冻结的 20 条评测集中，整体金标准行为符合 17/20（85%），"
        "三个不符合项为 ORCH-007 fan-in 来源完整性漏检、ORCH-016 失败传播状态不符，以及 ORCH-017 敏感薪资进入事件数据。"
        "这些负面结果均保留在分母中。",
    )
    replace_paragraph(
        document,
        "上述实验以 Git 提交 2c9c684 为代码基线",
        "上述实验以 Git 提交 2c9c684 为代码基线，于 2026 年 8 月 11 日执行。统一评测集包含 20 个唯一案例，"
        "SHA-256 为 19b59db4734ef206ef84ef49dcb325d0d39edae280a267145504cf82fe5b3ba3；MCP 盲测集包含 36 个唯一案例，"
        "SHA-256 为 7e091fab7b860288489257ef5e9c4bb54910c1c4acd36d8e6bc6016eab8bcb19。最终聚焦回归为 119 项通过、3 项预期失败，"
        "所有预期失败均作为不符合计入分母。计划门禁时延为本机独占进程内测量，不代表真实 HTTP 或生产墙钟时延。"
        "2026 年 8 月 12 日完成 MCP 选择器规则优化后，冻结盲测集保持不变，MCP 相关聚焦回归为 31 项通过。",
    )

    document.save(OUTPUT)


if __name__ == "__main__":
    main()
