from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
REPORT = ROOT / "doc-yz" / "多智能体协同与执行编排机制调研及原型成果报告-修订版-量化实验更新.docx"
OUTPUT = ROOT / "doc-yz" / "多智能体协同与执行编排机制调研及原型成果报告-修订版-量化实验更新-指标表述修正.docx"
NOTE = (
    "注：全场景自动闭环率以每种策略的全部 500 次试验为分母，其中包含不可重试、缺少可信备份及"
    "必须进入人工对账的场景，不等同于系统的一般任务成功率。全场景首次失败恢复率以发生首次失败的"
    "试验为分母。副作用结果不确定时进入 NEEDS_RECONCILIATION 属于预期的安全处置结果。"
)


def main() -> None:
    document = Document(REPORT)
    recovery_table = next(
        table
        for table in document.tables
        if table.rows[0].cells[0].text.strip() == "策略"
        and any("P95 虚拟时延" in cell.text for cell in table.rows[0].cells)
    )
    recovery_table.rows[0].cells[1].text = "全场景自动闭环率"
    recovery_table.rows[0].cells[2].text = "全场景首次失败恢复率"

    if not any(paragraph.text.strip() == NOTE for paragraph in document.paragraphs):
        analysis = next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.strip().startswith("实验覆盖 5 类故障")
        )
        note = document.add_paragraph(NOTE, style="Body Text")
        note._p.getparent().remove(note._p)
        analysis._p.addprevious(note._p)

    document.save(OUTPUT)


if __name__ == "__main__":
    main()
